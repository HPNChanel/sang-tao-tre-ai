"""Small typed data-plan interpreter. No generated source code is executed."""
from __future__ import annotations

from collections import OrderedDict
import datetime as dt
from decimal import Decimal, ROUND_HALF_UP
import json
import re

from openpyxl.utils import column_index_from_string, get_column_letter, range_boundaries

from .workbook import WorkbookError, normalized, number

MAX_ROWS = 20000
MAX_COLUMNS = 32
OPS = {'read','cell','filter','derive','aggregate','sort','take','project','union','unpivot','join'}
FIELDS = {
 'read': {'sheet','range','columns','fill','label'}, 'cell': {'sheet','cell'},
 'filter': {'input','where'}, 'derive': {'input','columns'},
 'aggregate': {'input','group_by','metrics'}, 'sort': {'input','by'},
 'take': {'input','count','offset','ties'}, 'project': {'input','columns'},
 'union': {'inputs'}, 'unpivot': {'input','columns','key','value'},
 'join': {'inputs','on','how','prefix'},
}


class PlanError(WorkbookError):
    pass


def integer(value, low, high, name):
    if type(value) is not int or not low <= value <= high:
        raise PlanError(f'{name} must be an integer in {low}..{high}')
    return value


def parse_date(value, locale='iso'):
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    if not isinstance(value,str):
        raise PlanError('date value unavailable')
    patterns = ['%Y-%m-%d','%Y/%m/%d']
    if locale in {'vi','dmy'}: patterns += ['%d/%m/%Y','%d-%m-%Y','%d.%m.%Y']
    if locale in {'en','mdy'}: patterns += ['%m/%d/%Y','%m-%d-%Y']
    for pattern in patterns:
        try:
            return dt.datetime.strptime(value.strip(),pattern).date()
        except ValueError:
            pass
    raise PlanError('unrecognized date; specify locale')


def numeric(value, locale='en'):
    result = number(value,locale)
    if result is None:
        raise PlanError('numeric operand unavailable')
    return result


def expression(node, row, depth=0):
    if depth > 8:
        raise PlanError('expression nesting limit')
    if not isinstance(node,dict):
        raise PlanError('expression must be an object')
    if node.get('op') == 'if' and set(node) == {'op','where','then','else'}:
        branch = 'then' if predicate(node['where'], row, depth+1) else 'else'
        return expression(node[branch], row, depth+1)
    if set(node)=={'col'}:
        if node['col'] not in row:
            raise PlanError(f'unknown expression column {node["col"]!r}; this step has {_visible(row)}')
        return row[node['col']]
    if set(node)=={'literal'}:
        value = node['literal']
        if value is None or isinstance(value,(str,bool)):
            return value
        if type(value) in (int,float):
            return numeric(value)
        raise PlanError('invalid literal')
    if set(node)-{'op','args','locale'} or node.get('op') is None:
        raise PlanError('invalid expression fields')
    args = node.get('args',[])
    if not isinstance(args,list) or not 1 <= len(args) <= 4:
        raise PlanError('expression needs 1..4 arguments')
    vals = [expression(arg,row,depth+1) for arg in args]
    op, locale = node['op'],node.get('locale','en')
    if op=='number' and len(vals)==1: return numeric(vals[0],locale)
    if op=='date' and len(vals)==1: return parse_date(vals[0],locale)
    if op=='text' and len(vals)==1: return '' if vals[0] is None else str(vals[0])
    if op=='concat': return ''.join('' if v is None else str(v) for v in vals)
    if op=='coalesce': return next((v for v in vals if v is not None and v!=''),None)
    if op in {'years','days','months'} and len(vals)==2:
        a,b = (parse_date(v,locale) for v in vals)
        if op=='days': return Decimal((b-a).days)
        if op=='months': return Decimal((b.year-a.year)*12+b.month-a.month-(b.day<a.day))
        return Decimal(b.year-a.year-((b.month,b.day)<(a.month,a.day)))
    if op=='year' and len(vals)==1: return Decimal(parse_date(vals[0],locale).year)
    if op=='month' and len(vals)==1: return parse_date(vals[0],locale).strftime('%Y-%m')
    numbers = [numeric(v,locale) for v in vals]
    if op=='abs' and len(numbers)==1: return abs(numbers[0])
    if op=='neg' and len(numbers)==1: return -numbers[0]
    if op=='round' and len(numbers)==2:
        digits = int(numbers[1])
        if numbers[1] != digits or not -12 <= digits <= 12:
            raise PlanError('invalid rounding precision')
        return numbers[0].quantize(Decimal(1).scaleb(-digits),rounding=ROUND_HALF_UP)
    if len(numbers)!=2:
        raise PlanError('arithmetic needs two arguments')
    a,b = numbers
    if op=='add': return a+b
    if op=='sub': return a-b
    if op=='mul': return a*b
    if op=='div' and b!=0: return a/b
    if op=='change_ratio' and a!=0: return (b-a)/a
    raise PlanError('unsupported operation or zero denominator')


def predicate(spec,row,depth=0):
    if depth>8 or not isinstance(spec,dict):
        raise PlanError('invalid filter')
    if set(spec) in ({'all'},{'any'}):
        key = next(iter(spec))
        items = spec[key]
        if not isinstance(items,list) or not 1<=len(items)<=12:
            raise PlanError('invalid filter group')
        # Validate every branch even when its truth value would short circuit.
        values = [predicate(item,row,depth+1) for item in items]
        return all(values) if key=='all' else any(values)
    if set(spec)-{'column','cmp','value','locale','normalize'}:
        raise PlanError('unknown filter field')
    if spec.get('column') not in row:
        raise PlanError(f'unknown filter column {spec.get("column")!r}; this step has {_visible(row)}')
    a, op, b = row[spec['column']],spec.get('cmp'),spec.get('value')
    if op=='is_null': return a is None or a==''
    if op=='not_null': return a is not None and a!=''
    if a is None:
        return False
    if op in {'in','not_in'}:
        if not isinstance(b,list) or len(b)>100:
            raise PlanError('membership requires a bounded list')
        found = any(predicate(dict(spec,cmp='eq',value=item),row,depth+1) for item in b)
        return found if op=='in' else not found
    if spec.get('normalize'):
        a,b = normalized(a),normalized(b)
    elif isinstance(a,Decimal):
        b = numeric(b,spec.get('locale','en'))
    elif isinstance(a,dt.date):
        a,b = parse_date(a),parse_date(b,spec.get('locale','iso'))
    if op=='eq': return a==b
    if op=='ne': return a!=b
    if op=='contains': return str(b) in str(a)
    if op=='not_contains': return str(b) not in str(a)
    if op=='starts_with': return str(a).startswith(str(b))
    try:
        if op=='lt': return a<b
        if op=='le': return a<=b
        if op=='gt': return a>b
        if op=='ge': return a>=b
    except TypeError:
        raise PlanError('filter types differ; derive a typed column first') from None
    raise PlanError('unsupported filter comparator')


def _visible(row):
    return ', '.join(sorted(k for k in row if not k.startswith('_')))[:200] or 'no columns'


def _bounded(rows):
    if len(rows)>MAX_ROWS:
        raise PlanError('record limit exceeded; narrow source range')
    return rows


def _names(columns):
    if not isinstance(columns,list) or not columns or len(columns)>MAX_COLUMNS or any(not isinstance(c,str) for c in columns):
        raise PlanError('invalid column list')
    return columns


def validate(plan):
    if not isinstance(plan,dict) or set(plan)-{'version','action','steps','result','format'}:
        raise PlanError('invalid CellPlanV2 fields')
    if plan.get('version')!=2 or plan.get('action')!='plan':
        raise PlanError('expected version 2 plan')
    steps = plan.get('steps')
    if not isinstance(steps,list) or not 1<=len(steps)<=16:
        raise PlanError('plan needs 1..16 steps')
    depths, readers, joins = {},0,0
    for step in steps:
        if not isinstance(step,dict) or step.get('op') not in OPS:
            raise PlanError('unknown step operation')
        sid,op = step.get('id'),step['op']
        if not isinstance(sid,str) or not re.fullmatch(r'[a-zA-Z][a-zA-Z0-9_]{0,31}',sid) or sid in depths:
            raise PlanError('step IDs must be unique simple names')
        if set(step)-({'id','op'}|FIELDS[op]):
            raise PlanError(f'unexpected fields for {op}')
        inputs = [] if op in {'read','cell'} else (step.get('inputs') if op in {'union','join'} else [step.get('input')])
        if not isinstance(inputs,list) or any(not isinstance(i,str) or i not in depths for i in inputs):
            raise PlanError(f'step {sid!r} inputs must name earlier step ids; {op} takes '+
                            ('inputs:[id,id]' if op in {'union','join'} else 'one input id; combine two steps with join'))
        if op=='join' and len(inputs)!=2: raise PlanError('join needs two inputs')
        if op=='union' and not 1<=len(inputs)<=3: raise PlanError('union needs 1..3 inputs')
        depth = 1+max((depths[i] for i in inputs),default=0)
        if depth>8: raise PlanError('plan depth limit')
        depths[sid] = depth
        readers += op in {'read','cell'}
        joins += op=='join'
    if readers>3 or joins>1 or plan.get('result') not in depths:
        raise PlanError('source/join/result limit')
    if not isinstance(plan.get('format',{}),dict):
        raise PlanError('format must be an object')
    return plan


def execute(plan,book):
    validate(plan)
    values, traces, sorted_by = {},[],{}
    for step in plan['steps']:
        book.check()
        op,sid = step['op'],step['id']
        source = values.get(step.get('input'))
        if op=='cell':
            item = book.metadata(step.get('sheet'),step.get('cell'))
            if item is None:
                raise PlanError('source cell does not exist')
            value = book.value(step['sheet'],step['cell'])
            if value is None:
                raise PlanError('source value unavailable')
            rows = [{'value':value,'_sheet':book.sheet_id(step['sheet']),'_row':book.coord(step['cell'])[0]}]
            trace = {'sheet':step['sheet'],'cell':step['cell']}
        elif op=='read':
            columns = step.get('columns')
            if not isinstance(columns,dict) or not 1<=len(columns)<=MAX_COLUMNS or any(not isinstance(k,str) or k.startswith('_') for k in columns):
                raise PlanError('read columns must map names to Excel column letters')
            try:
                c1,r1,c2,r2 = range_boundaries(step.get('range',''))
                positions = {k:column_index_from_string(v) for k,v in columns.items()}
            except (ValueError,TypeError,AttributeError):
                raise PlanError('invalid read range/columns') from None
            if None in (r1,c1,r2,c2) or r1<1 or c1<1 or r2<r1 or c2<c1 or r2>1048576 or c2>16384:
                raise PlanError('invalid read bounds')
            # The range bounds rows; explicit column letters may lie outside its columns.
            if r2-r1+1>MAX_ROWS:
                raise PlanError('read exceeds the record limit')
            fill = step.get('fill',[])
            if not isinstance(fill,list) or any(k not in columns for k in fill):
                raise PlanError('fill needs mapped columns')
            label = step.get('label',next(iter(columns)))
            if label not in columns: raise PlanError('unknown label column')
            sheet = book.sheet_id(step.get('sheet'))
            previous,stack,rows = {},[],[]
            for r in range(r1,r2+1):
                if r%100==0: book.check()
                row = {k:book.value(sheet,get_column_letter(c)+str(r),merged=True) for k,c in positions.items()}
                if all(v is None or v=='' for v in row.values()):
                    previous.clear(); stack.clear()
                    continue
                for k in fill:
                    if row[k] is None or row[k]=='': row[k] = previous.get(k)
                    else: previous[k] = row[k]
                meta = book.metadata(sheet,get_column_letter(positions[label])+str(r))
                indent = meta['indent'] if meta else 0
                if isinstance(row[label],str):
                    indent = max(indent,(len(row[label])-len(row[label].lstrip(' ')))//2)
                while stack and stack[-1][0]>=indent:
                    stack.pop()
                row.update(_row=Decimal(r),_sheet=Decimal(sheet),_indent=Decimal(indent),
                           _hidden=bool(meta and meta['hidden']),
                           _parent=stack[-1][1] if stack else '',
                           _path=' / '.join([str(x[1]) for x in stack]+[str(row[label] or '')]))
                if row[label] is not None: stack.append((indent,row[label]))
                rows.append(row)
            trace = {'sheet':sheet,'range':step['range'],'columns':columns}
        elif op=='filter':
            rows = [r for r in source if predicate(step.get('where'),r)]
            trace = {'input':step['input'],'where':step.get('where')}
        elif op=='derive':
            columns = step.get('columns')
            if not isinstance(columns,dict) or not 1<=len(columns)<=MAX_COLUMNS:
                raise PlanError('derive needs columns mapping')
            rows = []
            for row in source:
                updated = dict(row)
                for name,expr in columns.items():
                    if not isinstance(name,str) or name.startswith('_'):
                        raise PlanError('invalid derived column name')
                    updated[name] = expression(expr,updated)
                rows.append(updated)
            trace = {'input':step['input'],'columns':columns}
        elif op=='aggregate':
            by = step.get('group_by',[])
            if not isinstance(by,list) or len(by)>8 or any(not isinstance(k,str) for k in by):
                raise PlanError('invalid grouping')
            metrics = step.get('metrics')
            if not isinstance(metrics,dict) or not 1<=len(metrics)<=MAX_COLUMNS:
                raise PlanError('aggregate needs named metrics')
            groups = OrderedDict()
            for row in source:
                if any(k not in row for k in by): raise PlanError('unknown grouping column')
                key = tuple(row[k] for k in by)
                groups.setdefault(key,[]).append(row)
            if not groups and not by: groups[()] = []
            rows = []
            for key,records in groups.items():
                out = dict(zip(by,key))
                for name,metric in metrics.items():
                    if not isinstance(metric,dict) or set(metric)-{'fn','column','locale','where','weight'}:
                        raise PlanError('invalid metric fields')
                    fn,col,weight,locale = metric.get('fn'),metric.get('column'),metric.get('weight'),metric.get('locale','en')
                    for c in (col,weight):
                        if c is not None and any(c not in r for r in records):
                            raise PlanError(f'unknown metric column {c!r}; input has {_visible(records[0])}')
                    if (fn=='weighted_mean') != (weight is not None) or (fn=='weighted_mean' and col is None):
                        raise PlanError('weighted_mean needs column (values) and weight; weight is only for weighted_mean')
                    # A metric-level where is a conditional aggregate (SUMIFS-style) over this group.
                    chosen = [r for r in records if 'where' not in metric or predicate(metric['where'],r)]
                    raw = [r[col] for r in chosen if r[col] is not None and r[col]!=''] if col is not None else []
                    if fn=='count': out[name] = Decimal(len(raw) if col is not None else len(chosen))
                    elif fn=='count_distinct': out[name] = Decimal(len(set(raw)))
                    elif fn=='weighted_mean':
                        pairs = [(numeric(r[col],locale),numeric(r[weight],locale)) for r in chosen
                                 if r[col] not in (None,'') and r[weight] not in (None,'')]
                        total = sum((w for _,w in pairs),Decimal(0))
                        if not pairs or total==0: raise PlanError('weighted_mean has no observations or zero total weight')
                        out[name] = sum((v*w for v,w in pairs),Decimal(0))/total
                    else:
                        nums = [numeric(v,locale) for v in raw]
                        if not nums: raise PlanError('aggregate has no numeric observations')
                        if fn=='sum': out[name] = sum(nums,Decimal(0))
                        elif fn in {'mean','average'}: out[name] = sum(nums,Decimal(0))/len(nums)
                        elif fn=='min': out[name] = min(nums)
                        elif fn=='max': out[name] = max(nums)
                        elif fn=='median':
                            nums.sort(); n=len(nums)
                            out[name] = nums[n//2] if n%2 else (nums[n//2-1]+nums[n//2])/2
                        else: raise PlanError('unsupported metric')
                rows.append(out)
            trace = {'input':step['input'],'group_by':by,'metrics':metrics}
        elif op=='sort':
            by = step.get('by')
            if not isinstance(by,list) or not 1<=len(by)<=8:
                raise PlanError('sort needs 1..8 keys')
            rows = list(source)
            for item in reversed(by):
                if not isinstance(item,dict) or set(item)-{'column','descending','type','locale'} or not isinstance(item.get('column'),str):
                    raise PlanError('invalid sort key')
                col = item['column']
                if any(col not in row or row[col] is None for row in rows):
                    raise PlanError('sort key missing; filter nulls first')
                if type(item.get('descending',False)) is not bool: raise PlanError('descending must be boolean')
                kind = item.get('type','auto')
                if kind not in {'auto','number','text','date'}: raise PlanError('unsupported sort type')
                if kind=='auto' and any(isinstance(r[col],str) and number(r[col]) is not None for r in rows):
                    raise PlanError('numeric-looking text needs explicit sort type number or text')
                def sort_value(row):
                    if kind=='number': return numeric(row[col],item.get('locale','en'))
                    if kind=='text': return str(row[col])
                    if kind=='date': return parse_date(row[col],item.get('locale','iso'))
                    return row[col]
                try: rows.sort(key=sort_value,reverse=item.get('descending',False))
                except TypeError: raise PlanError('mixed sort types; derive a typed column first') from None
            sorted_by[sid] = [item['column'] for item in by]
            trace = {'input':step['input'],'by':by}
        elif op=='take':
            count = integer(step.get('count',1),1,MAX_ROWS,'count')
            offset = integer(step.get('offset',0),0,MAX_ROWS,'offset')
            rows = source[offset:offset+count]
            if step.get('ties',False):
                keys = sorted_by.get(step['input'])
                if not keys or not rows: raise PlanError('ties require a directly preceding sorted nonempty input')
                last = tuple(rows[-1][k] for k in keys)
                rows += [row for row in source[offset+count:] if tuple(row[k] for k in keys)==last]
            trace = {'input':step['input'],'offset':offset,'count':count,'ties':step.get('ties',False)}
        elif op=='project':
            columns = _names(step.get('columns'))
            if any(k not in r for r in source for k in columns): raise PlanError('unknown projection column')
            rows = [{k:r[k] for k in columns} for r in source]
            trace = {'input':step['input'],'columns':columns}
        elif op=='union':
            parts = [values[i] for i in step['inputs']]
            if sum(len(part) for part in parts)>MAX_ROWS: raise PlanError('union record limit')
            schemas = {tuple(sorted(k for k in r if not k.startswith('_'))) for part in parts for r in part}
            if len(schemas)>1: raise PlanError('union schemas differ')
            rows = _bounded([dict(r) for part in parts for r in part])
            trace = {'inputs':step['inputs']}
        elif op=='unpivot':
            columns = step.get('columns')
            if not isinstance(columns,dict) or not 1<=len(columns)<=MAX_COLUMNS:
                raise PlanError('unpivot columns map source names to labels')
            if len(source)*len(columns)>MAX_ROWS: raise PlanError('unpivot record limit')
            key,value = step.get('key','key'),step.get('value','value')
            if not isinstance(key,str) or not isinstance(value,str) or key==value:
                raise PlanError('invalid unpivot names')
            rows = []
            for row in source:
                if any(c not in row for c in columns): raise PlanError('unpivot column missing')
                base = {k:v for k,v in row.items() if k not in columns}
                if key in base or value in base: raise PlanError('unpivot names collide')
                for col,label in columns.items(): rows.append(dict(base,**{key:label,value:row[col]}))
            _bounded(rows)
            trace = {'input':step['input'],'columns':columns,'key':key,'value':value}
        elif op=='join':
            left,right = (values[i] for i in step['inputs'])
            how,prefix = step.get('how','inner'),step.get('prefix','r_')
            if how not in {'inner','left','cross'} or not isinstance(prefix,str) or len(prefix)>16:
                raise PlanError('invalid join mode/prefix')
            on = step.get('on',{} if how=='cross' else None)
            if how=='cross':
                if on!={}: raise PlanError('cross join takes no on keys')
                if len(left)*len(right)>MAX_ROWS: raise PlanError('cross join record limit')
            elif not isinstance(on,dict) or not 1<=len(on)<=8:
                raise PlanError('join on maps left aliases to right aliases, e.g. {"code":"item_code"}')
            right_columns = {k for row in right for k in row if not k.startswith('_')}
            left_columns = {k for row in left for k in row if not k.startswith('_')}
            unknown = [k for k in on if left and k not in left_columns]+[k for k in on.values() if right and k not in right_columns]
            if unknown:
                raise PlanError(f'unknown join key {unknown[0]!r}; left has {", ".join(sorted(left_columns))[:150]}; '
                                f'right has {", ".join(sorted(right_columns))[:150]}')
            def join_key(row,keys):
                found = tuple(row.get(k) for k in keys)
                if any(v is None or v=='' for v in found): return None
                return tuple(v.strip() if isinstance(v,str) else v for v in found)
            lookup = {}
            for row in (right if how!='cross' else []):
                key = join_key(row,list(on.values()))
                if key is None: continue  # blank keys never match: titles, gaps, notes
                if key in lookup: raise PlanError('right join key is not unique; use composite keys')
                lookup[key] = row
            rows = []
            for row in left:
                if how=='cross':
                    matches = right
                else:
                    match = lookup.get(join_key(row,list(on)))
                    if match is None and how=='inner': continue
                    matches = [match]
                for match in matches:
                    added = {prefix+k:(match.get(k) if match else None) for k in right_columns}
                    if set(row)&set(added): raise PlanError('join column collision; choose another prefix')
                    # Plain right aliases are also kept wherever the left side lacks that name.
                    plain = {k:added[prefix+k] for k in right_columns if k not in row and k not in added}
                    rows.append(dict(row,**added,**plain))
            trace = {'inputs':step['inputs'],'on':on,'how':how,'prefix':prefix}
        else:
            raise PlanError('unsupported operation')
        values[sid] = _bounded(rows)
        traces.append(dict(id=sid,op=op,rows=len(rows),**trace))
    result = values[plan['result']]
    if not result:
        raise PlanError('no matching evidence')
    return format_answer(result,plan.get('format',{})),traces


def format_answer(rows,fmt):
    allowed = {'columns','decimals','percent','scale','suffix','prefix','separator','row_separator','date_format','trim_zeros'}
    if set(fmt)-allowed:
        raise PlanError('unsupported answer format')
    columns = fmt.get('columns') or [k for k in rows[0] if not k.startswith('_')]
    _names(columns)
    if len(rows)>200 or len(columns)>8:
        raise PlanError('answer too large; project/select the requested result')
    if any(k not in r for r in rows for k in columns): raise PlanError('answer column missing')
    decimals = fmt.get('decimals')
    if decimals is not None: integer(decimals,0,12,'decimals')
    for key in ('suffix','prefix','separator','row_separator'):
        if key in fmt and (not isinstance(fmt[key],str) or len(fmt[key])>32):
            raise PlanError('invalid output delimiter')
    if type(fmt.get('percent',False)) is not bool or type(fmt.get('trim_zeros',True)) is not bool:
        raise PlanError('format flags must be boolean')
    scale = numeric(fmt.get('scale',1))
    if abs(scale)>Decimal('1e12'): raise PlanError('format scale too large')
    def render(value):
        if value is None: raise PlanError('answer value unavailable')
        if isinstance(value,Decimal):
            value *= scale * (100 if fmt.get('percent') else 1)
            if decimals is not None: value = value.quantize(Decimal(1).scaleb(-decimals),rounding=ROUND_HALF_UP)
            if value==0: value = abs(value)
            text = format(value,'f')
            if '.' in text and fmt.get('trim_zeros',decimals is None): text = text.rstrip('0').rstrip('.')
            return text+('%' if fmt.get('percent') else '')
        if isinstance(value,dt.date):
            pattern = {'iso':'%Y-%m-%d','dmy':'%d/%m/%Y','mdy':'%m/%d/%Y','year':'%Y'}.get(fmt.get('date_format','iso'))
            if pattern is None: raise PlanError('unsupported date format')
            return value.strftime(pattern)
        if isinstance(value,bool): return 'true' if value else 'false'
        return str(value).strip()
    return fmt.get('prefix','')+fmt.get('row_separator','; ').join(fmt.get('separator',', ').join(render(row[k]) for k in columns) for row in rows)+fmt.get('suffix','')
