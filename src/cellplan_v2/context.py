"""Question-specific cell views with explicit omission and coordinate provenance."""
from __future__ import annotations

import datetime as dt
from decimal import Decimal
import json
import re

from openpyxl.utils import get_column_letter

from .workbook import normalized, WorkbookError


def compact(value):
    return json.dumps(value, ensure_ascii=False, separators=(',', ':'))


def display(book, item):
    value = book.value(item['sheet'], item['cell'])
    if isinstance(value, Decimal):
        # Strings preserve decimal precision across JSON implementations.
        value = {'n': str(value)}
    elif isinstance(value, (dt.date, dt.datetime)):
        value = {'date': value.isoformat()[:10]}
    fields = {}
    if item['formula']:
        fields['f'] = item['formula']
    if item['format'] != 'General':
        fields['fmt'] = item['format']
    if item['indent']:
        fields['indent'] = item['indent']
    if item['hidden']:
        fields['hidden'] = True
    if fields:
        return dict(v=value, **fields)
    return value


def row_view(book, sid, r):
    rows = {}
    for item in book.cells(sid, f'A{r}:XFD{r}'):
        rows[re.sub(r'\d', '', item['cell'])] = display(book, item)
    return {'r': r, **rows}


def structure_hints(book):
    """Geometric clues only. All ranges and samples retain original coordinates."""
    hints = []
    for sheet in book.sheets:
        sid = sheet['id']
        groups = []
        start = last = None
        for (r,) in book.db.execute('SELECT DISTINCT r FROM cells WHERE sheet=? ORDER BY r', (sid,)):
            if last is not None and r > last + 1:
                groups.append((start, last))
                start = None
            if start is None:
                start = r
            last = r
        if last is not None:
            groups.append((start, last))
        for lo, hi in groups[:32]:
            c1, c2 = book.db.execute('SELECT min(c),max(c) FROM cells WHERE sheet=? AND r BETWEEN ? AND ?', (sid, lo, hi)).fetchone()
            samples = []
            for r in range(lo, min(hi, lo + 3) + 1):
                cells = list(book.cells(sid, f'{get_column_letter(c1)}{r}:{get_column_letter(c2)}{r}'))
                text = [x['cell'] for x in cells if x['type'] in {'s', 'str', 'inlineStr'}]
                nums = [x['cell'] for x in cells if x['type'] == 'n']
                samples.append({'row': r, 'text_cells': text[:24], 'numeric_cells': nums[:24]})
            hints.append({'sheet': sid, 'occupied_region': f'{get_column_letter(c1)}{lo}:{get_column_letter(c2)}{hi}',
                          'header_candidates': samples})
    return hints


def build_context(book, query, *, dual=False, max_bytes=15000, regions=None):
    """Full data when small; otherwise retain headers/tails/hits and name omissions.

    The executable index contains every cell, including rows not shown here.
    Selection never changes a source range or invents a cell value.
    """
    summary = book.summary()
    result = {'sheets': summary, 'complete': False, 'views': []}
    if dual:
        result['structure_hints_untrusted'] = structure_hints(book)
    base_bytes = len(compact(result).encode('utf-8'))
    if base_bytes >= max_bytes // 2:
        result.pop('structure_hints_untrusted', None)
        for item in summary:
            item['merges'] = item.get('merges', [])[:16]
        base_bytes = len(compact(result).encode('utf-8'))
    selected = {}
    if regions is not None:
        if not isinstance(regions, list) or not 1 <= len(regions) <= 3:
            raise WorkbookError('inspect requires 1..3 source regions')
        for region in regions:
            if not isinstance(region, dict) or set(region) != {'sheet', 'range'}:
                raise WorkbookError('inspect region fields: sheet, range')
            sid = book.sheet_id(region['sheet'])
            # Iteration validates bounds. Bound accumulation before materialization.
            rows = selected.setdefault(sid, set())
            for item in book.cells(sid, region['range']):
                rows.add(book.coord(item['cell'])[0])
                if len(rows) > 1000:
                    raise WorkbookError('inspect range has over 1000 occupied rows')
        priority = [(sid, r) for sid, rows in selected.items() for r in sorted(rows)]
    else:
        total = sum(s['cells'] for s in book.sheets)
        if total <= 3000:
            priority = [(s['id'], r) for s in book.sheets
                        for (r,) in book.db.execute('SELECT DISTINCT r FROM cells WHERE sheet=? ORDER BY r', (s['id'],))]
        else:
            priority = []
        # If the full context does not fit, a second pass selects useful rows.
    views, seen, used = {}, set(), base_bytes
    def include(sid, r):
        nonlocal used
        if (sid, r) in seen:
            return True
        book.check()
        view = row_view(book, sid, r)
        size = len(compact(view).encode('utf-8')) + 2
        if used + size > max_bytes - 256:
            return False
        views.setdefault(sid, []).append(view)
        seen.add((sid, r))
        used += size
        return True
    fits = bool(priority)
    for sid, r in priority:
        if not include(sid, r):
            fits = False
            break
    if regions is None and not fits:
        views, seen, used = {}, set(), base_bytes
        priority = []
        tokens = sorted(set(re.findall(r'[\w]+', normalized(query))), key=len, reverse=True)
        tokens = [t for t in tokens if len(t) >= 3][:18]
        for s in book.sheets:
            sid = s['id']
            bounds = s['bounds']
            if bounds[0] is None:
                continue
            first, last = bounds[0], bounds[2]
            priority.extend((sid, r) for r in range(first, min(first + 9, last + 1)))
            # Include separated block starts, not just the first table's header.
            prev = None
            for (r,) in book.db.execute('SELECT DISTINCT r FROM cells WHERE sheet=? ORDER BY r', (sid,)):
                if prev is not None and r > prev + 1:
                    priority.extend((sid, x) for x in range(r, min(r + 3, last + 1)))
                prev = r
            priority.extend((sid, r) for r in range(max(first, last - 2), last + 1))
        for sid, r in priority:
            include(sid, r)
        if tokens:
            clauses = ' OR '.join('instr(search,?)>0' for _ in tokens)
            for sid, r in book.db.execute(f'SELECT DISTINCT sheet,r FROM cells WHERE {clauses} ORDER BY sheet,r LIMIT 1500', tokens):
                include(sid, r)
        for s in book.sheets:
            for (r,) in book.db.execute('SELECT DISTINCT r FROM cells WHERE sheet=? ORDER BY r LIMIT 300', (s['id'],)):
                include(s['id'], r)
    total_rows = book.db.execute('SELECT count(*) FROM (SELECT sheet,r FROM cells GROUP BY sheet,r)').fetchone()[0]
    result['complete'] = len(seen) == total_rows
    for sid, rows in views.items():
        rows.sort(key=lambda x: x['r'])
        result['views'].append({'sheet': sid, 'rows': rows})
    result['omitted_rows'] = total_rows - len(seen)
    result['reading_note'] = ('All stored cells are shown.' if result['complete'] else
                              'Rows are omitted. Execute read over the full required range; do not aggregate shown samples. Inspect missing headers if needed.')
    text = compact(result)
    if len(text.encode('utf-8')) > max_bytes:
        raise WorkbookError('context metadata exceeds request budget')
    return result
