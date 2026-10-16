"""Streaming OOXML index; source coordinates are never flattened away.

SQLite is scratch space, not a data dependency or an answer cache. Shared
strings and cell values stay on disk. No external workbook links are followed.
"""
from __future__ import annotations

import ast
import datetime as dt
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
import hashlib
import json
from pathlib import Path, PurePosixPath
import re
import sqlite3
import tempfile
import time
import unicodedata
import xml.etree.ElementTree as ET
from zipfile import ZipFile

from openpyxl.styles.numbers import BUILTIN_FORMATS, is_date_format
from openpyxl.utils import column_index_from_string, get_column_letter, range_boundaries
from openpyxl.utils.datetime import from_excel, CALENDAR_MAC_1904, CALENDAR_WINDOWS_1900

NS = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
REL = '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}'
MAX_CELL_TEXT = 262144
MAX_INDEX_BYTES = 48 * 1024 * 1024


class WorkbookError(ValueError):
    pass


def normalized(value):
    text = unicodedata.normalize('NFKD', str(value)).casefold().replace('đ', 'd')
    return ' '.join(''.join(c for c in text if not unicodedata.combining(c)).split())


def number(value, locale='en'):
    if value is None or isinstance(value, (bool, dt.date, dt.datetime)):
        return None
    if isinstance(value, Decimal):
        return value if value.is_finite() else None
    text = str(value).strip().replace('\u00a0', '').replace(' ', '')
    if not text or text.casefold() in {'na', 'n/a', 'none', 'null', '-', '—', '…'}:
        return None
    negative = text.startswith('(') and text.endswith(')')
    if negative:
        text = text[1:-1]
    percent = text.endswith('%')
    text = text.rstrip('%').lstrip('$€£₫')
    if locale == 'vi':
        text = text.replace('.', '').replace(',', '.')
    else:
        text = text.replace(',', '')
    try:
        result = Decimal(text)
        if not result.is_finite():
            return None
        return (-result if negative else result) / (100 if percent else 1)
    except InvalidOperation:
        return None


class WorkbookIndex:
    def __init__(self, path, deadline=None):
        self.path = Path(path).resolve()
        self.deadline = deadline
        self.temp = tempfile.TemporaryDirectory(prefix='cellqa-')
        self.db_path = Path(self.temp.name) / 'cells.sqlite'
        self.db = sqlite3.connect(self.db_path)
        self.db.execute('PRAGMA journal_mode=OFF')
        self.db.execute('PRAGMA cache_size=-8192')
        self.db.execute('PRAGMA temp_store=FILE')
        self.db.executescript('''
          CREATE TABLE strings(id INTEGER PRIMARY KEY,value TEXT);
          CREATE TABLE cells(sheet INTEGER,r INTEGER,c INTEGER,value TEXT,kind TEXT,
            formula TEXT,format TEXT,indent INTEGER,hidden INTEGER,search TEXT,
            PRIMARY KEY(sheet,r,c)) WITHOUT ROWID;
          CREATE TABLE merges(sheet INTEGER,r1 INTEGER,c1 INTEGER,r2 INTEGER,c2 INTEGER);
        ''')
        self.sheets = []
        self.epoch = CALENDAR_WINDOWS_1900
        self.formula_cache = {}
        self.formula_stack = set()
        self.digest = hashlib.sha256()
        with self.path.open('rb') as handle:
            for chunk in iter(lambda: handle.read(1024 * 1024), b''):
                self.check()
                self.digest.update(chunk)
        self.sha256 = self.digest.hexdigest()
        try:
            self._load()
        except BaseException:
            self.close()
            raise

    def check(self):
        if self.deadline is not None and time.monotonic() >= self.deadline:
            raise WorkbookError('workbook deadline exceeded')
        if self.db_path.exists() and self.db_path.stat().st_size > MAX_INDEX_BYTES:
            raise WorkbookError('workbook index disk budget exceeded')

    def close(self):
        if getattr(self, 'db', None) is not None:
            self.db.close()
            self.db = None
        if getattr(self, 'temp', None) is not None:
            self.temp.cleanup()
            self.temp = None

    def _load(self):
        with ZipFile(self.path) as book:
            workbook = ET.fromstring(book.read('xl/workbook.xml'))
            props = workbook.find(NS + 'workbookPr')
            if props is not None and props.get('date1904') in {'1', 'true'}:
                self.epoch = CALENDAR_MAC_1904
            rels = ET.fromstring(book.read('xl/_rels/workbook.xml.rels'))
            targets = {}
            for entry in rels:
                if entry.get('TargetMode') == 'External':
                    continue
                target = entry.get('Target', '')
                name = target.lstrip('/') if target.startswith('/') else 'xl/' + target
                if '..' not in PurePosixPath(name).parts:
                    targets[entry.get('Id')] = name
            formats, styles = {}, []
            if 'xl/styles.xml' in book.namelist():
                root = ET.fromstring(book.read('xl/styles.xml'))
                for item in root.findall(NS + 'numFmts/' + NS + 'numFmt'):
                    formats[int(item.get('numFmtId'))] = item.get('formatCode', '')
                for item in root.findall(NS + 'cellXfs/' + NS + 'xf'):
                    fmt = int(item.get('numFmtId', '0'))
                    align = item.find(NS + 'alignment')
                    styles.append((formats.get(fmt, BUILTIN_FORMATS.get(fmt, 'General')),
                                   int(float(align.get('indent', '0'))) if align is not None else 0))
            if not styles:
                styles = [('General', 0)]
            if 'xl/sharedStrings.xml' in book.namelist():
                with book.open('xl/sharedStrings.xml') as stream:
                    index = 0
                    stack = []
                    for event, element in ET.iterparse(stream, events=('start', 'end')):
                        if event == 'start':
                            stack.append(element)
                        else:
                            if element.tag == NS + 'si':
                                value = ''.join(t.text or '' for t in element.iter(NS + 't'))
                                if len(value) > MAX_CELL_TEXT:
                                    raise WorkbookError('oversized shared string')
                                self.db.execute('INSERT INTO strings VALUES(?,?)', (index, value))
                                index += 1
                                if len(stack) > 1:
                                    stack[-2].remove(element)
                                if index % 1000 == 0:
                                    self.db.commit()
                                    self.check()
                            stack.pop()
            for sheet in workbook.findall(NS + 'sheets/' + NS + 'sheet'):
                target = targets.get(sheet.get(REL + 'id'))
                if not target or not target.startswith('xl/worksheets/'):
                    continue
                sid = len(self.sheets)
                info = {'id': sid, 'name': sheet.get('name'), 'state': sheet.get('state', 'visible')}
                self.sheets.append(info)
                row_hidden = 0
                stack = []
                with book.open(target) as stream:
                    count = 0
                    for event, element in ET.iterparse(stream, events=('start', 'end')):
                        if event == 'start':
                            stack.append(element)
                            if element.tag == NS + 'row':
                                row_hidden = int(element.get('hidden', '0') in {'1', 'true'})
                            continue
                        if element.tag == NS + 'c':
                            coord = element.get('r', '')
                            match = re.fullmatch(r'([A-Z]{1,3})([1-9][0-9]{0,6})', coord)
                            if not match:
                                raise WorkbookError('invalid cell coordinate')
                            r, c = int(match[2]), column_index_from_string(match[1])
                            if r > 1048576 or c > 16384:
                                raise WorkbookError('cell outside Excel bounds')
                            kind = element.get('t', 'n')
                            vn, fn = element.find(NS + 'v'), element.find(NS + 'f')
                            value = vn.text if vn is not None else None
                            if kind == 's' and value is not None:
                                found = self.db.execute('SELECT value FROM strings WHERE id=?', (int(value),)).fetchone()
                                value = found[0] if found else None
                                kind = 's'
                            elif kind == 'inlineStr':
                                value = ''.join(t.text or '' for t in element.iter(NS + 't'))
                                kind = 's'
                            formula = fn.text if fn is not None else None
                            if fn is not None and not formula:
                                formula = '#UNSUPPORTED_SHARED_FORMULA'
                            if value is not None and len(value) > MAX_CELL_TEXT:
                                raise WorkbookError('oversized cell')
                            style = int(element.get('s', '0'))
                            fmt, indent = styles[style] if 0 <= style < len(styles) else styles[0]
                            if value is not None or formula is not None:
                                self.db.execute('INSERT INTO cells VALUES(?,?,?,?,?,?,?,?,?,?)',
                                                (sid,r,c,value,kind,formula,fmt,indent,row_hidden,normalized(value or '')))
                            count += 1
                            element.clear()
                        elif element.tag == NS + 'mergeCell':
                            c1,r1,c2,r2 = range_boundaries(element.get('ref'))
                            self.db.execute('INSERT INTO merges VALUES(?,?,?,?,?)', (sid,r1,c1,r2,c2))
                        elif element.tag == NS + 'row':
                            if len(stack) > 1:
                                stack[-2].remove(element)
                            self.check()
                            if count % 1000 < 164:
                                self.db.commit()
                                if self.db_path.stat().st_size > MAX_INDEX_BYTES:
                                    raise WorkbookError('workbook index disk budget exceeded')
                        stack.pop()
                bounds = self.db.execute('SELECT min(r),min(c),max(r),max(c),count(*) FROM cells WHERE sheet=?', (sid,)).fetchone()
                info.update(bounds=list(bounds[:4]), cells=bounds[4])
        self.db.commit()
        self.check()

    def sheet_id(self, sheet):
        if isinstance(sheet, int) and not isinstance(sheet, bool) and 0 <= sheet < len(self.sheets):
            return sheet
        for item in self.sheets:
            if item['name'] == sheet:
                return item['id']
        raise WorkbookError('unknown sheet')

    def coord(self, ref):
        if not isinstance(ref, str) or not re.fullmatch(r'\$?[A-Z]{1,3}\$?[1-9][0-9]{0,6}', ref):
            raise WorkbookError('invalid cell reference')
        match = re.fullmatch(r'\$?([A-Z]{1,3})\$?([0-9]+)', ref)
        r, c = int(match[2]), column_index_from_string(match[1])
        if r > 1048576 or c > 16384:
            raise WorkbookError('cell outside Excel bounds')
        return r,c

    def metadata(self, sheet, ref):
        sid = self.sheet_id(sheet)
        r,c = self.coord(ref)
        item = self.db.execute('SELECT value,kind,formula,format,indent,hidden FROM cells WHERE sheet=? AND r=? AND c=?', (sid,r,c)).fetchone()
        if item is None:
            return None
        value,kind,formula,fmt,indent,hidden = item
        return {'sheet':sid,'cell':ref.replace('$',''),'value':value,'type':kind,
                'formula':formula,'format':fmt,'indent':indent,'hidden':bool(hidden)}

    def value(self, sheet, ref, merged=False):
        sid = self.sheet_id(sheet)
        item = self.metadata(sid, ref)
        if item is None and merged:
            r,c = self.coord(ref)
            merge = self.db.execute('SELECT r1,c1 FROM merges WHERE sheet=? AND r1<=? AND r2>=? AND c1<=? AND c2>=? LIMIT 1', (sid,r,r,c,c)).fetchone()
            if merge:
                item = self.metadata(sid, get_column_letter(merge[1])+str(merge[0]))
                # Merge expansion is only a label hint. Numeric cells are
                # physically present once, so expansion must never double them.
                if item is not None and item['type'] not in {'s','str','inlineStr'}:
                    return None
        if item is None:
            return None
        raw = item['value']
        if raw is None and item['formula']:
            return self._formula(sid, item['cell'], item['formula'])
        if item['type'] == 'n' and raw is not None:
            try:
                n = Decimal(raw)
                if not n.is_finite():
                    return None
                if is_date_format(item['format']):
                    return from_excel(float(n), self.epoch)
                return n
            except (InvalidOperation, ValueError, OverflowError):
                return None
        if item['type'] == 'd' and raw:
            try:
                return dt.datetime.fromisoformat(raw)
            except ValueError:
                return None
        if item['type'] == 'b':
            return raw == '1'
        if item['type'] == 'e':
            return None
        return raw

    def _formula(self, sid, coord, formula):
        key = (sid, coord)
        if key in self.formula_cache:
            return self.formula_cache[key]
        if key in self.formula_stack or len(self.formula_stack) >= 32 or len(formula) > 4096:
            return None
        self.formula_stack.add(key)
        try:
            refs = {}
            pattern = r"(?:(?:'[^']+'|[A-Za-z_][A-Za-z0-9_ ]*)!)?\$?[A-Z]{1,3}\$?[1-9][0-9]*(?::\$?[A-Z]{1,3}\$?[1-9][0-9]*)?"
            def substitute(match):
                text = match[0]
                sheet, address = (text.rsplit('!',1) if '!' in text else (sid,text))
                if isinstance(sheet,str):
                    sheet = sheet.strip("'").replace("''", "'")
                address = address.replace('$','')
                if ':' in address:
                    c1,r1,c2,r2 = range_boundaries(address)
                    if (r2-r1+1)*(c2-c1+1) > 100000:
                        raise WorkbookError('formula range too large')
                    vals = []
                    for r in range(r1,r2+1):
                        self.check()
                        for c in range(c1,c2+1):
                            reference = get_column_letter(c)+str(r)
                            value = self.value(sheet,reference)
                            meta = self.metadata(sheet,reference)
                            if value is None and meta is not None and (meta['formula'] or meta['type']=='e'):
                                raise WorkbookError('formula range contains unavailable operand')
                            vals.append(value)
                    val = [v for v in vals if isinstance(v,Decimal)]
                else:
                    val = self.value(sheet,address)
                    meta = self.metadata(sheet,address)
                    if val is None and meta is not None and (meta['formula'] or meta['type']=='e'):
                        raise WorkbookError('formula operand unavailable')
                name = f'v_{len(refs)}'
                refs[name] = val
                return name
            expression = re.sub(pattern, substitute, formula.lstrip('=')).replace('^','**').replace(';',',')
            tree = ast.parse(expression, mode='eval')
            if len(list(ast.walk(tree))) > 256:
                raise WorkbookError('formula too complex')
            def calculate(node, depth=0):
                if depth > 24:
                    raise WorkbookError('formula nesting limit')
                if isinstance(node,ast.Expression):
                    return calculate(node.body,depth+1)
                if isinstance(node,ast.Constant) and type(node.value) in (int,float):
                    return Decimal(ast.get_source_segment(expression,node))
                if isinstance(node,ast.Name) and node.id in refs:
                    return refs[node.id]
                if isinstance(node,ast.UnaryOp) and isinstance(node.op,(ast.UAdd,ast.USub)):
                    v = calculate(node.operand,depth+1)
                    return -v if isinstance(node.op,ast.USub) else v
                if isinstance(node,ast.BinOp):
                    a,b = calculate(node.left,depth+1),calculate(node.right,depth+1)
                    # Excel treats truly empty cell references as zero in
                    # arithmetic; failed formulas/errors remain unavailable.
                    a = Decimal(0) if a is None else a
                    b = Decimal(0) if b is None else b
                    if not isinstance(a,Decimal) or not isinstance(b,Decimal):
                        raise WorkbookError('formula operands must be numeric')
                    if isinstance(node.op,ast.Add): return a+b
                    if isinstance(node.op,ast.Sub): return a-b
                    if isinstance(node.op,ast.Mult): return a*b
                    if isinstance(node.op,ast.Div): return a/b
                    if isinstance(node.op,ast.Pow) and b == int(b) and abs(b) <= 16: return a**int(b)
                if isinstance(node,ast.Call) and isinstance(node.func,ast.Name) and not node.keywords:
                    args = [calculate(a,depth+1) for a in node.args]
                    flat = [v for a in args for v in (a if isinstance(a,list) else [a])]
                    flat = [v for v in flat if isinstance(v,Decimal)]
                    fn = node.func.id.upper()
                    if fn=='SUM': return sum(flat,Decimal(0))
                    if fn=='COUNT': return Decimal(len(flat))
                    if fn=='AVERAGE' and flat: return sum(flat,Decimal(0))/len(flat)
                    if fn=='MIN' and flat: return min(flat)
                    if fn=='MAX' and flat: return max(flat)
                    if fn=='ABS' and len(args)==1: return abs(args[0])
                    if fn=='ROUND' and len(args) in (1,2):
                        places = int(args[1]) if len(args)==2 else 0
                        if abs(places) <= 12:
                            return args[0].quantize(Decimal(1).scaleb(-places),rounding=ROUND_HALF_UP)
                raise WorkbookError('unsupported formula')
            result = calculate(tree)
            result = result if isinstance(result,Decimal) and result.is_finite() else None
        except (ValueError,TypeError,ArithmeticError,KeyError,SyntaxError,RecursionError):
            result = None
        finally:
            self.formula_stack.discard(key)
        if len(self.formula_cache) < 8192:
            self.formula_cache[key] = result
        return result

    def cells(self, sheet, region=None):
        sid = self.sheet_id(sheet)
        bounds = self.sheets[sid]['bounds']
        if not bounds[0]:
            return
        r1,c1,r2,c2 = bounds
        if region is not None:
            try:
                c1,r1,c2,r2 = range_boundaries(region.replace('$',''))
            except (ValueError,AttributeError):
                raise WorkbookError('invalid region') from None
            if None in (r1,c1,r2,c2) or r1<1 or c1<1 or r2>1048576 or c2>16384:
                raise WorkbookError('invalid region bounds')
        cursor = self.db.execute('SELECT r,c,value,kind,formula,format,indent,hidden FROM cells WHERE sheet=? AND r BETWEEN ? AND ? AND c BETWEEN ? AND ? ORDER BY r,c', (sid,r1,r2,c1,c2))
        for row in cursor:
            r,c,value,kind,formula,fmt,indent,hidden = row
            yield {'sheet':sid,'cell':get_column_letter(c)+str(r),'value':value,'type':kind,
                   'formula':formula,'format':fmt,'indent':indent,'hidden':bool(hidden)}

    def summary(self):
        out = []
        for sheet in self.sheets:
            item = dict(sheet)
            item['merges'] = [f'{get_column_letter(c1)}{r1}:{get_column_letter(c2)}{r2}'
                              for r1,c1,r2,c2 in self.db.execute('SELECT r1,c1,r2,c2 FROM merges WHERE sheet=? LIMIT 128',(sheet['id'],))]
            out.append(item)
        return out
