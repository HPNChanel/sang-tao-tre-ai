import datetime as dt
import json
import os
import sys

import openpyxl
import pytest

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))
from src.code_planner import plan
from src.plan_executor import execute
from src.xlsx_tables import load_tables

SCHEMA = json.load(open(os.path.join(os.path.dirname(__file__), "..", "src", "plan_schema.json")))


def _ask(path, q):
    t = load_tables(path)
    p = plan(q, t, os.path.basename(path))
    assert p and set(p) <= set(SCHEMA["properties"]) and p["op"] in SCHEMA["properties"]["op"]["enum"]
    return execute(p, t, q)


@pytest.fixture
def sales(tmp_path):
    wb = openpyxl.Workbook(); ws = wb.active
    ws.append(["Sales (units)"]); ws.append([])
    ws.append(["Region", "Jan", "Feb", "Mar", "Total"])
    for r in [("North", 10, 20, 30), ("South", 5, 5, 50), ("East", 7, 8, 9)]:
        ws.append(list(r) + [f"=SUM(B{ws.max_row + 1}:D{ws.max_row + 1})"])
    ws.append(["Grand total", "=SUM(B4:B6)", "=SUM(C4:C6)", "=SUM(D4:D6)", "=SUM(E4:E6)"])
    p = tmp_path / "s.xlsx"; wb.save(p); return str(p)


def test_lookup_sum_range_argmax_compare(sales):
    assert _ask(sales, "What were Feb sales in South? Answer with only the integer.") == "5"
    assert _ask(sales, "What were total Jan-Mar sales for North? Answer with only the integer.") == "60"
    assert _ask(sales, "Which region had the highest sales in Mar? Answer with only the region name.") == "South"
    assert _ask(sales, "Were Jan sales in North greater than Jan sales in East? Answer with only Yes or No.") == "Yes"
    assert _ask(sales, "What is the sum of Jan sales across all regions? Answer with only the integer.") == "22"
    assert _ask(sales, "What was the average Feb sales? Answer with only the number rounded to 2 decimal places.") == "11.00"


def test_dates_count_nth(tmp_path):
    wb = openpyxl.Workbook(); ws = wb.active
    ws.append(["Name", "Birth", "Admission", "Score"])
    for n, b, a, s in [("An", 2000, 2020, 50), ("Binh", 1990, 2024, 70), ("Chi", 1995, 2015, 90)]:
        ws.append([n, dt.datetime(b, 1, 1), dt.datetime(a, 1, 1), s])
    p = tmp_path / "d.xlsx"; wb.save(p)
    # age at admission: An 20, Binh 34, Chi 20 -> oldest Binh
    assert _ask(str(p), "What score did the oldest person receive? Answer with only the integer score.") == "70"
    assert _ask(str(p), "How many people scored at least 60? Answer with only the integer.") == "2"
    assert _ask(str(p), "On what date was the earliest admission? Answer in YYYY-MM-DD format.") == "2015-01-01"


def test_hidden_sheet_skipped(tmp_path):
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Now"
    ws.append(["Item", "Stock"]); ws.append(["A", 1]); ws.append(["B", 2])
    ws.row_dimensions[3].hidden = True
    h = wb.create_sheet("Old"); h.append(["Item", "Stock"]); h.append(["A", 100]); h.sheet_state = "hidden"
    p = tmp_path / "h.xlsx"; wb.save(p)
    assert _ask(str(p), "What is the total stock across all items? Answer with only the integer.") == "3"
