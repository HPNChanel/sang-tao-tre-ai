import sys
from pathlib import Path

import openpyxl
import pytest

sys.path.insert(0, str(Path(__file__).parents[1]))
from src.xlsx_tables import load_tables


@pytest.mark.parametrize("offset", [0, 3, 7])
@pytest.mark.parametrize("years", [["2020", "2021"], [2020, 2021]])
def test_year_and_text_headers_after_spacer(tmp_path, offset, years):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.cell(1, offset + 1, "A survey")
    rows = [["District", "Class", *years, "Visits"],
            ["Pine", "Coastal", 20, 30, 7], ["Oak", "Inland", 40, 55, 9]]
    for r, row in enumerate(rows, 4):
        for c, value in enumerate(row, offset + 1):
            ws.cell(r, c, value)
    path = tmp_path / "survey.xlsx"
    wb.save(path)
    tables = load_tables(path)
    assert len(tables) == 1
    table = tables[0]
    assert table.header_rows == [4]
    assert table.label_col == "District"
    assert table.df["_row_path"].tolist() == ["Pine", "Oak"]
    assert table.df["2021"].tolist() == [30, 55]
    assert table.df["Visits"].sum() == 16


def test_real_year_valued_body_row_is_not_a_header(tmp_path):
    wb = openpyxl.Workbook()
    for row in [["District", "Opened", "Closed"], ["Pine", 1980, 2021], ["Oak", 1990, 2024]]:
        wb.active.append(row)
    path = tmp_path / "history.xlsx"
    wb.save(path)
    table = load_tables(path)[0]
    assert table.header_rows == [1]
    assert table.df["_row_path"].tolist() == ["Pine", "Oak"]
