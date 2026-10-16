import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from src.planner import build_views, match_rows
from src.textmatch import Query
from src.xlsx_tables import load_tables


def test_named_parent_outweighs_its_shorter_child(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Item", "Weight"])
    ws.append(["Grains and seeds", 40])
    ws.append(["Grains", 30])
    ws.cell(3, 1).alignment = Alignment(indent=1)
    ws.append(["Textiles", 60])
    path = tmp_path / "products.xlsx"
    wb.save(path)
    view = build_views(load_tables(path))[0]
    rows, _ = match_rows(view, Query("Compare Grains and seeds with Textiles"))
    assert [x[0] for x in rows] == ["Grains and seeds", "Textiles"]


def test_ancestor_words_cannot_consume_other_leaf_operand(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Item", "Cost"])
    ws.append(["Cloth and timber"])
    ws.append(["Cloth", 3])
    ws.append(["Timber", 7])
    path = tmp_path / "materials.xlsx"
    wb.save(path)
    view = build_views(load_tables(path))[0]
    rows, _ = match_rows(view, Query("Cloth minus Timber"))
    assert [x[0] for x in rows] == ["Cloth and timber > Cloth", "Cloth and timber > Timber"]


def test_exact_label_can_share_tokens_with_column(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Item", "Fair value"])
    ws.append(["Total assets at fair value", 20])
    ws.append(["Other", 10])
    path = tmp_path / "assets.xlsx"
    wb.save(path)
    view = build_views(load_tables(path))[0]
    q = Query("Total assets at fair value")
    rows, _ = match_rows(view, q, frozenset(q.match("Fair value")[1]))
    assert len(rows) == 1
    assert rows[0][0] == "Total assets at fair value"
