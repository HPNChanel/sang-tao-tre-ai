"""Vietnamese row filters on a table built in the test.

A cue that names one column and one value restricts the aggregate.
A cue that cannot be tied to one column is left blank.
"""
import os
import sys

import openpyxl

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))
from src.code_planner import plan, plan_with_confidence
from src.plan_executor import execute
from src.xlsx_tables import load_tables


PEOPLE = [
    {"name": "An", "shift": "day", "city": "H1", "team": "red", "area": "north",
     "score": 10, "bonus": 1, "desk": "day"},
    {"name": "Bo", "shift": "night", "city": "H2", "team": "red", "area": "south",
     "score": 20, "bonus": 5, "desk": "day"},
    {"name": "Cy", "shift": "night", "city": "H1", "team": "blue", "area": "south",
     "score": 30, "bonus": 5, "desk": "day"},
    {"name": "Di", "shift": "evening", "city": "H2", "team": "north", "area": "east",
     "score": 40, "bonus": 9, "desk": "day"},
    {"name": "Ed", "shift": "night", "city": "H2", "team": "red", "area": "east",
     "score": 50, "bonus": 5, "desk": "night"},
]


def _score(pred):
    return sum(row["score"] for row in PEOPLE if pred(row))


def _book(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Name", "Shift", "City", "Team", "Area", "Score", "Bonus", "Desk"])
    for row in PEOPLE:
        ws.append([row["name"], row["shift"], row["city"], row["team"], row["area"],
                   row["score"], row["bonus"], row["desk"]])
    path = tmp_path / "shifts.xlsx"
    wb.save(path)
    return str(path)


def _ask(path, query):
    tables = load_tables(path)
    planned, conf, reasons = plan_with_confidence(query, tables)
    gated = plan(query, tables, os.path.basename(path))
    if gated is None:
        return None, conf, reasons, planned
    return execute(gated, tables, query), conf, reasons, gated


def test_unique_filters_sum_the_group(tmp_path):
    path = _book(tmp_path)
    night_h2 = _score(lambda r: r["shift"] == "night" and r["city"] == "H2")
    day_h1 = _score(lambda r: r["shift"] == "day" and r["city"] == "H1")
    evening_h2 = _score(lambda r: r["shift"] == "evening" and r["city"] == "H2")
    evening_h1 = _score(lambda r: r["shift"] == "evening" and r["city"] == "H1")
    night_only = _score(lambda r: r["shift"] == "night")
    city_h2 = _score(lambda r: r["city"] == "H2")
    below_30 = _score(lambda r: r["score"] < 30)
    night_h1 = _score(lambda r: r["shift"] == "night" and r["city"] == "H1")
    whole = _score(lambda r: True)
    assert night_h2 != whole and night_h2 == 70

    cases = [
        ("Hãy chỉ xét người làm ca night tại thành phố H2. "
         "Tổng điểm Score của các dòng này là bao nhiêu? Chỉ trả lời bằng số nguyên.",
         str(night_h2)),
        ("Những người làm ca day ở thành phố H1, tổng điểm Score là bao nhiêu? "
         "Chỉ trả lời bằng số nguyên.",
         str(day_h1)),
        ("Thuộc thành phố H2 mà làm ca evening, tổng điểm Score là bao nhiêu? "
         "Chỉ trả lời bằng số nguyên.",
         str(evening_h2)),
        ("Những người làm ca evening tại thành phố H1, tổng điểm Score là bao nhiêu? "
         "Chỉ trả lời bằng số nguyên.",
         str(evening_h1)),
        ("Chỉ xét những người làm ca night, tổng điểm Score là bao nhiêu? "
         "Chỉ trả lời bằng số nguyên.",
         str(night_only)),
        ("Chỉ xét H2, tổng điểm Score là bao nhiêu? Chỉ trả lời bằng số nguyên.",
         str(city_h2)),
        ("Chỉ xét những bản ghi có điểm chưa tới 30, tổng điểm Score là bao nhiêu? "
         "Chỉ trả lời bằng số nguyên.",
         str(below_30)),
        ("Những người làm ca night ở thành phố H1, điểm Score cao nhất là bao nhiêu? "
         "Chỉ trả lời bằng số nguyên.",
         str(night_h1)),
        ("Hãy chỉ xét người làm ca tối tại thành phố H2. "
         "Tổng điểm Score của các dòng này là bao nhiêu? Chỉ trả lời bằng số nguyên.",
         str(night_h2)),
    ]
    for query, expected in cases:
        got, conf, reasons, planned = _ask(path, query)
        assert got == expected, (query, got, expected, conf, reasons, None if not planned else planned["rows"]["where"])
        assert conf >= 0.5
        assert got != str(whole)


def test_exact_threshold_counts_rows(tmp_path):
    path = _book(tmp_path)
    exact = sum(1 for row in PEOPLE if row["score"] == 40)
    got, conf, reasons, _planned = _ask(
        path, "Với cột Score, có bao nhiêu dòng đúng bằng 40? Ghi chữ số."
    )
    assert got == str(exact), (got, conf, reasons)
    assert exact == 1


def test_mean_of_named_shift(tmp_path):
    path = _book(tmp_path)
    vals = [row["score"] for row in PEOPLE if row["shift"] == "night"]
    # 20, 30, 50 -> 33.3 with half-up to 1 decimal
    assert vals == [20, 30, 50]
    got, conf, reasons, _planned = _ask(
        path,
        "Trung bình điểm Score của những người làm ca night là bao nhiêu? "
        "Chỉ trả lời bằng số, làm tròn đến 1 chữ số thập phân.",
    )
    assert got == "33.3", (got, conf, reasons)


def test_unresolved_filter_is_blank(tmp_path):
    path = _book(tmp_path)
    whole = str(_score(lambda r: True))
    queries = [
        "Chỉ xét những người làm ca swing, tổng điểm Score là bao nhiêu? Chỉ trả lời bằng số nguyên.",
        "Chỉ xét các dòng rồi cho biết tổng điểm Score là bao nhiêu? Chỉ trả lời bằng số nguyên.",
        "Chỉ xét những người thuộc north, tổng điểm Score là bao nhiêu? Chỉ trả lời bằng số nguyên.",
        "Chỉ xét night, tổng điểm Score là bao nhiêu? Chỉ trả lời bằng số nguyên.",
    ]
    scored, conf, reasons, _planned = _ask(
        path,
        "Những người có mức chưa tới 30, tổng điểm Score là bao nhiêu? Chỉ trả lời bằng số nguyên.",
    )
    assert scored == "30", (scored, conf, reasons)
    for query in queries:
        got, conf, _reasons, _planned = _ask(path, query)
        assert got is None, (query, got)
        assert conf == 0
        assert got != whole


def test_questions_without_a_row_filter_stay_on_the_whole_table(tmp_path):
    path = _book(tmp_path)
    whole = _score(lambda r: True)
    highest = max(row["score"] for row in PEOPLE)
    at_least_40 = sum(1 for row in PEOPLE if row["score"] >= 40)
    assert (_ask(path, "Tổng điểm Score của tất cả mọi người là bao nhiêu? Chỉ trả lời bằng số nguyên.")[0]
            == str(whole))
    assert (_ask(path, "Điểm Score cao nhất là bao nhiêu? Chỉ trả lời bằng số nguyên.")[0]
            == str(highest))
    assert (_ask(path, "Với cột Score, có bao nhiêu dòng từ 40 trở lên? Ghi chữ số.")[0]
            == str(at_least_40))
    assert (_ask(path, "What is the total Score? Answer with only the integer.")[0] == str(whole))
    assert (_ask(path, "How many people scored at least 40? Answer with only the integer.")[0]
            == str(at_least_40))


def test_filter_code_does_not_pin_a_one_letter_row(tmp_path):
    """A code splits into letters. The letter must not become the only row."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Name", "Shift", "City", "Score"])
    rows = [
        ("A", "day", "K4", 100),
        ("B", "night", "K7", 10),
        ("K", "night", "K4", 7),
        ("D", "evening", "K7", 4),
        ("E", "night", "K7", 20),
    ]
    for row in rows:
        ws.append(list(row))
    path = tmp_path / "codes.xlsx"
    wb.save(path)
    group = sum(score for _name, shift, city, score in rows if shift == "night" and city == "K7")
    whole = sum(score for *_rest, score in rows)
    got, conf, reasons, planned = _ask(
        str(path),
        "Những người làm ca night ở thành phố K7, tổng điểm Score là bao nhiêu? "
        "Chỉ trả lời bằng số nguyên.",
    )
    assert group == 30 and group != whole
    assert got == str(group), (got, conf, reasons, None if not planned else planned["rows"])
    assert planned["rows"]["paths"] == []


def test_number_words_and_row_labels(tmp_path):
    wb = openpyxl.Workbook()
    people = wb.active
    people.title = "People"
    people.append(["Ma", "Kip", "Kho", "KhoiLuong"])
    people.append(["An", "dem", "K4", 72])
    people.append(["Bo", "dem", "K4", 10])
    people.append(["Cy", "ngay", "K7", 72])
    sales = wb.create_sheet("Sales")
    sales.append(["Region", "Sales"])
    sales.append(["South", 12])
    sales.append(["North", 4])
    path = tmp_path / "mixed.xlsx"
    wb.save(path)
    counted, conf, reasons, _planned = _ask(
        str(path),
        "Kiện thuộc kip dem với KhoiLuong đúng bằng bảy mươi hai thì có mấy kiện? Viết một chữ số.",
    )
    assert counted == "1", (counted, conf, reasons)
    region, conf, reasons, planned = _ask(
        str(path),
        "Doanh số Sales của vùng South là bao nhiêu? Chỉ trả lời bằng số nguyên.",
    )
    assert region == "12", (region, conf, reasons, None if not planned else planned["rows"])


def test_vietnamese_headers_use_the_same_cues(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Tên", "Ca", "Thành phố", "Điểm"])
    rows = [("A", "night", "H2", 5), ("B", "night", "H1", 7), ("C", "day", "H2", 9)]
    for row in rows:
        ws.append(list(row))
    path = tmp_path / "ca.xlsx"
    wb.save(path)
    expected = sum(score for _name, shift, city, score in rows if shift == "night" and city == "H2")
    whole = sum(score for _name, _shift, _city, score in rows)
    got, conf, reasons, _planned = _ask(
        str(path),
        "Những người làm ca night ở thành phố H2, tổng điểm là bao nhiêu? Chỉ trả lời bằng số nguyên.",
    )
    assert got == str(expected), (got, expected, whole, conf, reasons)
    assert expected != whole


def _save(wb, tmp_path, name):
    path = tmp_path / name
    wb.save(path)
    return str(path)


def _people_banner(tmp_path):
    """Two header rows, the way a real sheet groups columns under one banner."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.merge_cells("B1:E1")
    ws["B1"] = "Record"
    for col, name in enumerate(["Name", "Gender", "City", "Shift", "Score"], 1):
        ws.cell(2, col, name)
    rows = [
        ("An", "nam", "H1", "day", 10),
        ("Bo", "nu", "H1", "night", 300),
        ("K", "nu", "K7", "night", 7),
        ("Cy", "nam", "K7", "night", 250),
        ("Di", "nam", "K7", "day", 250),
        ("Ed", "nam", "H1", "night", 40),
    ]
    for r, row in enumerate(rows, 3):
        for c, value in enumerate(row, 1):
            ws.cell(r, c, value)
    return _save(wb, tmp_path, "banner.xlsx"), rows


def test_gender_word_is_not_the_number_five():
    from src.vi import translate
    gender, _ = translate("Những người nam đạt đúng hai trăm điểm")
    assert "five" not in gender.lower()
    counted, _ = translate("Trong năm người lớn tuổi nhất")
    assert "five" in counted.lower()


def test_banner_subset_sum_ignores_partial_row(tmp_path):
    path, rows = _people_banner(tmp_path)
    group = sum(score for _n, gender, city, _s, score in rows if gender == "nam" and city == "K7")
    whole = sum(score for *_rest, score in rows)
    got, conf, reasons, planned = _ask(
        path,
        "Những người thuộc nam ở thành phố K7, tổng điểm Score là bao nhiêu? "
        "Chỉ trả lời bằng số nguyên.",
    )
    assert group == 500 and group != whole
    assert got == str(group), (got, conf, reasons, None if not planned else planned["rows"])
    assert "K" not in (planned["rows"]["paths"] if planned else [])


def test_banner_count_hundreds_keeps_the_group(tmp_path):
    path, rows = _people_banner(tmp_path)
    group = sum(1 for _n, gender, _c, _s, score in rows if gender == "nam" and score == 250)
    got, conf, reasons, planned = _ask(
        path,
        "Đếm dòng thuộc nam khi điểm Score đúng bằng hai trăm năm mươi. Ghi chữ số.",
    )
    assert group == 2
    assert got == "2", (got, conf, reasons, None if not planned else planned)
    assert planned["op"] == "count"


def test_banner_top_n_keeps_k_and_filters_the_group(tmp_path):
    path, _rows = _people_banner(tmp_path)
    got, conf, reasons, planned = _ask(
        path,
        "Top 3 người thuộc nam theo điểm Score. Liệt kê tên, phân cách bởi dấu phẩy và một dấu cách.",
    )
    assert planned and planned["args"].get("k") == 3, (planned, conf, reasons)
    assert got, (got, conf, reasons)
    names = [part.strip() for part in str(got).split(",")]
    assert names == ["Cy", "Di", "Ed"] or names == ["Di", "Cy", "Ed"], (got, planned["rows"])
    assert "Bo" not in names


def test_banner_month_span_uses_the_year(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.merge_cells("B1:E1")
    ws["B1"] = 2024
    ws.merge_cells("F1:I1")
    ws["F1"] = 2023
    ws["A2"] = "Region"
    for i, month in enumerate(["Mar", "Apr", "May", "Jun"]):
        ws.cell(2, 2 + i, month)
        ws.cell(2, 6 + i, month)
    ws["A3"] = "North"
    for i, value in enumerate([10, 20, 30, 40, 1, 2, 3, 4]):
        ws.cell(3, 2 + i, value)
    ws["A4"] = "South"
    for i, value in enumerate([7, 8, 9, 11, 5, 5, 5, 5]):
        ws.cell(4, 2 + i, value)
    path = _save(wb, tmp_path, "months.xlsx")
    got, conf, reasons, planned = _ask(
        path,
        "Tổng từ tháng 3 đến tháng 5 năm 2024 của vùng North là bao nhiêu? "
        "Chỉ trả lời bằng số nguyên.",
    )
    assert got == "60", (got, conf, reasons, None if not planned else planned["columns"])
    blank, conf, reasons, _planned = _ask(
        path,
        "Tổng từ tháng 3 đến tháng 5 năm 2025 của vùng North là bao nhiêu? "
        "Chỉ trả lời bằng số nguyên.",
    )
    assert blank is None, (blank, conf, reasons)


def test_banner_ambiguous_category_is_blank(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.merge_cells("B1:D1")
    ws["B1"] = "Record"
    for col, name in enumerate(["Name", "Team", "Desk", "Score"], 1):
        ws.cell(2, col, name)
    for r, row in enumerate([("An", "day", "day", 10), ("Bo", "night", "day", 20)], 3):
        for c, value in enumerate(row, 1):
            ws.cell(r, c, value)
    path = _save(wb, tmp_path, "ambig.xlsx")
    got, conf, reasons, _planned = _ask(
        path,
        "Những người thuộc day, tổng điểm Score là bao nhiêu? Chỉ trả lời bằng số nguyên.",
    )
    assert got is None, (got, conf, reasons)
    assert conf == 0.0


def test_greater_than_words_counts_with_the_shift(tmp_path):
    """A word threshold and a shift both apply. The shift alone is the wrong count."""
    path, rows = _people_banner(tmp_path)
    both = sum(1 for _n, _g, _c, shift, score in rows if shift == "night" and score > 200)
    shift_only = sum(1 for _n, _g, _c, shift, _s in rows if shift == "night")
    assert both == 2 and shift_only == 4
    got, conf, reasons, planned = _ask(
        path,
        "Đếm dòng làm ca night có điểm Score lớn hơn hai trăm. Ghi chữ số.",
    )
    assert got == "2", (got, conf, reasons, None if not planned else planned["rows"])
    assert planned["op"] == "count"
    assert got != str(shift_only)


def test_unbound_comparison_does_not_keep_the_shift(tmp_path):
    """Score and Bonus are both numeric. A bare 'lớn hơn' must not count the shift only."""
    path = _book(tmp_path)
    shift_only = sum(1 for row in PEOPLE if row["shift"] == "night")
    got, conf, reasons, planned = _ask(
        path,
        "Đếm dòng làm ca night lớn hơn hai trăm. Ghi chữ số.",
    )
    assert got is None, (got, conf, reasons, None if not planned else planned["rows"])
    assert conf == 0
    assert got != str(shift_only)


def test_over_thousands_in_a_year_counts_rows(tmp_path):
    """'vượt N nghìn' plus a year counts rows. It does not return one income."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.merge_cells("A1:D1")
    ws["A1"] = "Bang"
    for col, name in enumerate(["Province", "Group", "Year", "Thu nhập"], 1):
        ws.cell(2, col, name)
    data = [
        ("Ha Noi", "north", 2024, 6000),
        ("Ha Noi", "north", 2023, 1000),
        ("Hue", "north", 2024, 8000),
        ("Hue", "north", 2023, 9000),
        ("Da Nang", "south", 2024, 7000),
    ]
    for r, row in enumerate(data, 3):
        for c, value in enumerate(row, 1):
            ws.cell(r, c, value)
    path = _save(wb, tmp_path, "provinces.xlsx")
    got, conf, reasons, planned = _ask(
        path,
        "Đếm tỉnh thuộc north trong năm 2024 có thu nhập vượt 5 nghìn. Ghi chữ số.",
    )
    assert got == "2", (got, conf, reasons, None if not planned else (planned["op"], planned["rows"]))
    assert planned["op"] == "count"
    assert got not in {"6000", "8000", "9000", "7000"}


def test_vietnamese_month_headers_use_the_year(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.merge_cells("B1:E1")
    ws["B1"] = 2024
    ws.merge_cells("F1:I1")
    ws["F1"] = 2023
    ws["A2"] = "Region"
    for i, month in enumerate(["Tháng 3", "Tháng 4", "Tháng 5", "Tháng 6"]):
        ws.cell(2, 2 + i, month)
        ws.cell(2, 6 + i, month)
    ws["A3"] = "North"
    for i, value in enumerate([10, 20, 30, 40, 1, 2, 3, 4]):
        ws.cell(3, 2 + i, value)
    ws["A4"] = "South"
    for i, value in enumerate([7, 8, 9, 11, 5, 5, 5, 5]):
        ws.cell(4, 2 + i, value)
    path = _save(wb, tmp_path, "vn-months.xlsx")
    got, conf, reasons, planned = _ask(
        path,
        "Tổng từ tháng ba đến tháng năm năm 2024 của vùng North là bao nhiêu? "
        "Ghi chữ số.",
    )
    assert got == "60", (got, conf, reasons, None if not planned else planned["columns"])
    blank, conf, reasons, _planned = _ask(
        path,
        "Tổng từ tháng 3 đến tháng 5 năm 2025 của vùng North là bao nhiêu? Ghi chữ số.",
    )
    assert blank is None, (blank, conf, reasons)


def test_gender_section_ranks_inside_the_group(tmp_path):
    """nam/nữ are section rows. Ranking stays inside nam and does not read nam as five."""
    from src.vi import translate
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.merge_cells("B1:B1")
    ws["B1"] = "Record"
    ws["A2"] = "Name"
    ws["B2"] = "Score"
    ws["A3"] = "nam"
    ws["A4"] = "An"
    ws["B4"] = 10
    ws["A5"] = "Bo"
    ws["B5"] = 40
    ws["A6"] = "nu"
    ws["A7"] = "Cy"
    ws["B7"] = 90
    path = _save(wb, tmp_path, "sections.xlsx")
    text, _hits = translate("Top 1 người thuộc nam theo điểm Score.")
    assert "five" not in text.lower()
    got, conf, reasons, planned = _ask(
        path,
        "Top 2 người thuộc nam theo điểm Score. Liệt kê tên, phân cách bởi dấu phẩy và một dấu cách.",
    )
    assert planned and planned["args"].get("k") == 2, (planned, conf, reasons)
    names = [part.strip() for part in str(got).split(",")]
    assert names == ["Bo", "An"], (got, conf, reasons, None if not planned else planned["rows"])
    assert "Cy" not in names


def test_shift_word_does_not_pin_one_row(tmp_path):
    """A row named like the shift must not shrink a count that also has a threshold."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.merge_cells("B1:C1")
    ws["B1"] = "Record"
    for col, name in enumerate(["Ten", "Age", "Ca"], 1):
        ws.cell(2, col, name)
    rows = [
        ("sang", 55, "sang"),
        ("An", 60, "sang"),
        ("Bo", 51, "sang"),
        ("Cy", 40, "sang"),
        ("Di", 70, "toi"),
    ]
    for r, row in enumerate(rows, 3):
        for c, value in enumerate(row, 1):
            ws.cell(r, c, value)
    path = _save(wb, tmp_path, "crew.xlsx")
    got, conf, reasons, planned = _ask(
        path,
        "Những dòng làm ca sang mà tuổi lớn hơn năm mươi, đếm giúp mấy dòng? Ghi chữ số.",
    )
    assert got == "3", (got, conf, reasons, None if not planned else planned["rows"])
    assert planned["op"] == "count"
    assert "sang" not in [str(p).split(" > ")[-1].lower() for p in planned["rows"]["paths"]]


def test_age_band_row_is_not_blocked(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.merge_cells("A1:C1")
    ws["A1"] = "Bang"
    ws["A2"] = "Group"
    ws["B2"] = "Employed"
    ws["C2"] = "Hours"
    ws["A3"] = "55 years and over"
    ws["B3"] = 10
    ws["C3"] = 3
    ws["A4"] = "25 to 54"
    ws["B4"] = 4
    ws["C4"] = 9
    path = _save(wb, tmp_path, "bands.xlsx")
    got, conf, reasons, _planned = _ask(
        path,
        "Tổng Employed của nhóm từ 55 tuổi trở lên là bao nhiêu? Ghi chữ số.",
    )
    assert got == "10", (got, conf, reasons)


def test_women_section_phrase_is_not_blocked(tmp_path):
    """'nữ giới' is the women row. It must not be read as a missing column filter."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.merge_cells("A1:B1")
    ws["A1"] = "Bang"
    ws["A2"] = "Group"
    ws["B2"] = "Score"
    ws["A3"] = "Women"
    ws["B3"] = 8
    ws["A4"] = "Men"
    ws["B4"] = 30
    path = _save(wb, tmp_path, "groups.xlsx")
    got, conf, reasons, _planned = _ask(
        path,
        "Ở nhóm nữ giới, tổng điểm Score là bao nhiêu? Ghi chữ số.",
    )
    assert got == "8", (got, conf, reasons)


def test_age_column_band_still_filters(tmp_path):
    """'từ 40 tuổi trở lên' on a numeric age column is a real cutoff, not a row label."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Ten", "Tuoi", "Ca"])
    ws.append(["An", 60, "sang"])
    ws.append(["Bo", 40, "sang"])
    ws.append(["Cy", 30, "toi"])
    path = _save(wb, tmp_path, "ages.xlsx")
    got, conf, reasons, planned = _ask(
        path,
        "Có mấy người từ 40 tuổi trở lên? Ghi chữ số.",
    )
    assert got == "2", (got, conf, reasons, None if not planned else planned["rows"])
    assert planned["op"] == "count"


def test_shift_alias_row_does_not_shrink_the_count(tmp_path):
    """A row named like the shift must not hide the other rows that share its alias."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.merge_cells("A1:C1")
    ws["A1"] = "Kho"
    for col, name in enumerate(["Ma", "Gio", "Kip"], 1):
        ws.cell(2, col, name)
    rows = [("Dem", 25, "night"), ("An", 30, "night"), ("Bo", 22, "night"), ("Cy", 18, "night"), ("Di", 40, "day")]
    for r, row in enumerate(rows, 3):
        for c, value in enumerate(row, 1):
            ws.cell(r, c, value)
    path = _save(wb, tmp_path, "kips.xlsx")
    got, conf, reasons, planned = _ask(
        path,
        "Máy có Gio lớn hơn hai mươi thuộc kip dem, đếm mấy máy? Viết một chữ số.",
    )
    assert got == "3", (got, conf, reasons, None if not planned else planned["rows"])
    assert planned["op"] == "count"
    assert "dem" not in [str(p).split(" > ")[-1].lower() for p in planned["rows"]["paths"]]


def test_empty_filtered_count_is_blank(tmp_path):
    """A filter that keeps no row is blank, not a confident zero."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Ma", "Gio", "Kip"])
    ws.append(["An", 25, "night"])
    ws.append(["Bo", 30, "night"])
    path = _save(wb, tmp_path, "empty.xlsx")
    got, conf, reasons, _planned = _ask(
        path,
        "Đếm máy thuộc kip dem khi Gio đúng bằng chín trăm. Viết một chữ số.",
    )
    assert got is None, (got, conf, reasons)
    assert got != "0"


def test_men_phrase_keeps_a_selected_row(tmp_path):
    """'nam giới' is the men column. It must not be dropped as a missing value."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Group", "Men", "Women"])
    ws.append(["55 years and over", 30, 31])
    ws.append(["20 to 24 years", 20, 21])
    path = _save(wb, tmp_path, "genders.xlsx")
    got, conf, reasons, _planned = _ask(
        path,
        "Tổng cột Men thuộc nhóm nam giới. Ghi chữ số.",
    )
    assert got == "50", (got, conf, reasons)


def test_row_after_o_is_not_blocked(tmp_path):
    """ở + a row label is the row itself. A miss must not blank a correct lookup."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.merge_cells("A1:B1")
    ws["A1"] = "Bang"
    ws["A2"] = "Region"
    ws["B2"] = "Sales"
    ws["A3"] = "Hue"
    ws["B3"] = 12
    ws["A4"] = "North"
    ws["B4"] = 4
    path = _save(wb, tmp_path, "places.xlsx")
    got, conf, reasons, planned = _ask(
        path,
        "Doanh số Sales ở Hue là bao nhiêu? Ghi chữ số.",
    )
    assert got == "12", (got, conf, reasons, None if not planned else planned["rows"])


def test_function_word_rows_do_not_shrink_a_filtered_count(tmp_path):
    """Rows named like 'có' or 'làm' must not be the only rows a filtered count sees.

    Both the shift and the word threshold have to apply. The intersection is 3.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Ten", "Nang", "Ca"])
    rows = [
        ("Lam", 80, "sang"),
        ("Co", 90, "toi"),
        ("An", 75, "sang"),
        ("Bo", 73, "sang"),
        ("Cy", 20, "sang"),
        ("Di", 90, "toi"),
    ]
    for row in rows:
        ws.append(list(row))
    path = _save(wb, tmp_path, "lots.xlsx")
    got, conf, reasons, planned = _ask(
        path,
        "Có mấy kiện làm ca sang khi cột Nang vượt bảy mươi hai? Viết một chữ số.",
    )
    assert got == "3", (got, conf, reasons, None if not planned else planned["rows"])
    assert planned["op"] == "count"
    wheres = planned["rows"]["where"]
    assert len(wheres) == 2, wheres
    assert {item["op"] for item in wheres} == {"==", ">"}
    leaves = [str(p).split(" > ")[-1].lower() for p in planned["rows"]["paths"]]
    assert "lam" not in leaves and "co" not in leaves


def test_one_unbound_condition_does_not_count_one(tmp_path):
    """Two numeric columns and no named threshold: do not count the shift alone."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Ten", "Nang", "Phu", "Ca"])
    for row in [("Lam", 80, 1, "sang"), ("Co", 90, 2, "toi"), ("An", 75, 3, "sang"), ("Bo", 73, 4, "sang")]:
        ws.append(list(row))
    path = _save(wb, tmp_path, "two-num.xlsx")
    got, conf, reasons, _planned = _ask(
        path,
        "Hãy đếm máy thuộc ca sang vượt mức bảy mươi hai. Trả về một số nguyên.",
    )
    assert got is None, (got, conf, reasons)
    assert got != "1"
    assert conf == 0


def test_yes_no_keeps_the_base_comparison(tmp_path):
    """A shift cue that cannot bind must not blank a Yes/No comparison."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Region", "Sales"])
    ws.append(["K4", 40])
    ws.append(["K7", 12])
    path = _save(wb, tmp_path, "yesno.xlsx")
    got, conf, reasons, planned = _ask(
        path,
        "Doanh số Sales của K4 có lớn hơn K7 không? Làm ca sang. Chỉ trả lời Yes hoặc No.",
    )
    assert got == "Yes", (got, conf, reasons, None if not planned else planned["rows"])
    assert planned["op"] == "compare"
    assert conf >= 0.5


def test_value_list_keeps_both_numbers(tmp_path):
    """An unbound shift cue must not blank a list of two values."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Name", "Amount"])
    ws.append(["Bo", 1200])
    ws.append(["Cy", 3400])
    path = _save(wb, tmp_path, "pair.xlsx")
    got, conf, reasons, planned = _ask(
        path,
        "Liệt kê Amount của Bo và Cy, phân cách bởi dấu phẩy và một dấu cách. Thuộc kip dem.",
    )
    assert got == "1200, 3400", (got, conf, reasons, None if not planned else planned)
    assert conf >= 0.5
