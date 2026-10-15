"""Caption-path tables keep a simple lookup or an unfiltered total.

Count, a threshold then an aggregate, a two-cell pair, and a year-to-year
difference are blank. The sheet is a wharf tally, not a copied question.
"""
import os
import sys

import openpyxl

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))
from src.code_planner import plan_with_confidence
from src.plan_executor import execute
from src.xlsx_tables import load_tables


def _wharf(path, caption):
    wb = openpyxl.Workbook()
    ws = wb.active
    if caption:
        ws["B1"] = "Wharf tally card"
    row0 = 2 if caption else 1
    ws.cell(row0, 1, "Slip")
    ws.cell(row0, 2, 2016)
    ws.cell(row0, 3, 2017)
    for i, (name, a, b) in enumerate(
        [("Kelp", 4, 9), ("Moss", 11, 2), ("Reef", 7, 8)], start=row0 + 1
    ):
        ws.cell(i, 1, name)
        ws.cell(i, 2, a)
        ws.cell(i, 3, b)
    wb.save(path)
    return load_tables(str(path))


def _see(tables, query):
    planned, conf, reasons = plan_with_confidence(query, tables)
    return planned, conf, reasons


def test_caption_path_blanks_the_four_ops(tmp_path):
    tables = _wharf(tmp_path / "wharf.xlsx", caption=True)
    assert tables[0].title_skipped is True

    q = "What is the sum of 2017? Answer with only the integer."
    planned, conf, reasons = _see(tables, q)
    assert "blank:refined-header" not in reasons
    assert conf >= 0.5
    assert execute(planned, tables, q) == "19"

    q = "What is 2017 recorded against Moss?"
    planned, conf, reasons = _see(tables, q)
    assert "blank:refined-header" not in reasons
    assert conf >= 0.5
    assert execute(planned, tables, q) == "2"

    q = "How many slips have 2017 of 8 or more?"
    planned, conf, reasons = _see(tables, q)
    assert planned["op"] == "count"
    assert conf == 0.0
    assert reasons == ["blank:refined-header"] or "blank:refined-header" in reasons

    q = "What is the sum of 2017 for slips with 2016 greater than 6? Answer with only the integer."
    planned, conf, reasons = _see(tables, q)
    assert planned["op"] == "sum"
    assert planned["rows"]["where"]
    assert conf == 0.0
    assert "blank:refined-header" in reasons

    q = "2016 and 2017 for Kelp."
    planned, conf, reasons = _see(tables, q)
    assert planned["op"] == "values"
    assert conf == 0.0
    assert "blank:refined-header" in reasons

    q = "What is the difference between 2017 and 2016 for Kelp?"
    planned, conf, reasons = _see(tables, q)
    assert planned["op"] == "diff"
    assert conf == 0.0
    assert "blank:refined-header" in reasons


def test_flat_header_does_not_take_the_caption_blank(tmp_path):
    tables = _wharf(tmp_path / "flat.xlsx", caption=False)
    assert tables[0].title_skipped is False
    q = "How many slips have 2017 of 8 or more?"
    planned, conf, reasons = _see(tables, q)
    assert "blank:refined-header" not in reasons
    assert conf >= 0.5
    assert execute(planned, tables, q) == "2"
