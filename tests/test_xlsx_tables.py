import os, sys
import openpyxl
from openpyxl.styles import Alignment
import pytest
sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))
from src.code_planner import plan
from src.plan_executor import execute
from src.xlsx_tables import load_tables, summarize, to_number, clean_text, strip_footnotes


def _cols(t):
    return [c for c in t.df.columns if not str(c).startswith("_")]


def _values(t):
    return [c for c in _cols(t) if c != t.label_col]


def _ask(tables, q):
    pl = plan(q, tables)
    return execute(pl, tables, q), pl


def _save(wb, tmp_path, name="t.xlsx"):
    p = tmp_path / name
    wb.save(p)
    return str(p)


def test_to_number():
    assert to_number("1,234") == 1234
    assert to_number("12.5(1)") == 12.5
    assert to_number("-") is None and to_number("(NA)") is None
    assert to_number("(5)") == -5
    assert to_number(" 3 ") == 3


def test_clean_and_footnotes():
    assert clean_text("Con-\nstruction") == "Construction"
    assert clean_text("Self-\nemployed") == "Self-employed"
    assert strip_footnotes("Hospital services(1)(3)") == "Hospital services"
    assert strip_footnotes("Trade balance 3/") == "Trade balance"


def test_merged_headers_title_notes_sections(tmp_path):
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "S"
    ws["A1"] = "TABLE 1. Something long title"; ws.merge_cells("A1:E1")
    ws["A3"] = "Group"; ws.merge_cells("A3:A4")
    ws["B3"] = 2023; ws.merge_cells("B3:E3")
    ws["B4"] = "Men"; ws["C4"] = "Women"; ws["D4"] = "Total(1)"; ws["E4"] = "Rate"
    ws["A5"] = "SECTION A"
    ws.append(["Alpha", "1,000", 2, 3, "-"])
    ws.append(["Beta(2)", 4, "(NA)", 6, 1.5])
    ws["A9"] = "SECTION B"
    ws.append(["Alpha", 7, 8, 9, 2.5])
    ws["A12"] = "Footnotes:"
    ws["A13"] = "(1) Includes things."; ws.merge_cells("A13:E13")
    ts = load_tables(_save(wb, tmp_path))
    assert len(ts) == 1
    t = ts[0]
    assert "TABLE 1" in t.title
    assert t.meta == ["2023"]
    cols = [c for c in t.df.columns if not c.startswith("_")]
    assert cols == ["Group", "Men", "Women", "Total", "Rate"]
    assert t.df["_row_path"].tolist() == ["SECTION A > Alpha", "SECTION A > Beta", "SECTION B > Alpha"]
    assert t.df["Men"].tolist() == [1000, 4, 7]
    assert t.df["Women"].isna().tolist() == [False, True, False]
    assert len(t.notes) == 2
    assert "Men" in summarize(ts)


def test_two_level_merged_header_and_indent(tmp_path):
    wb = openpyxl.Workbook(); ws = wb.active
    ws["A1"] = "Item"; ws.merge_cells("A1:A2")
    ws["B1"] = "Index"; ws.merge_cells("B1:C1")
    ws["B2"] = "Sep"; ws["C2"] = "Oct"
    rows = [("All", 0, 10, 11), ("Food", 1, 5, 6), ("Bread", 2, 2, 3), ("Energy", 1, 4, 4)]
    for i, (n, ind, a, b) in enumerate(rows, start=3):
        ws.cell(i, 1, n).alignment = Alignment(indent=ind)
        ws.cell(i, 2, a); ws.cell(i, 3, b)
    t = load_tables(_save(wb, tmp_path))[0]
    assert [c for c in t.df.columns if not c.startswith("_")] == ["Item", "Index | Sep", "Index | Oct"]
    assert t.df["_row_path"].tolist() == ["All", "All > Food", "All > Food > Bread", "All > Energy"]


def test_leading_space_indent(tmp_path):
    wb = openpyxl.Workbook(); ws = wb.active
    ws.append(["Name", "Val"])
    for n, v in [("Total", 3), ("  A", 1), ("  B", 2)]:
        ws.append([n, v])
    t = load_tables(_save(wb, tmp_path))[0]
    assert t.df["_row_path"].tolist() == ["Total", "Total > A", "Total > B"]


def test_multiple_tables_one_sheet(tmp_path):
    wb = openpyxl.Workbook(); ws = wb.active
    ws["B1"] = "Fiscal"; ws.merge_cells("B1:C1"); ws["B2"] = 2023; ws["C2"] = 2024
    ws["A3"] = "Exports"; ws["B3"] = 1.5; ws["C3"] = 2.5
    ws["A4"] = "Imports"; ws["B4"] = 3; ws["C4"] = 4
    ws["B6"] = "Calendar"; ws.merge_cells("B6:C6"); ws["B7"] = 2023; ws["C7"] = 2024
    ws["A8"] = "Exports"; ws["B8"] = 9; ws["C8"] = 8
    ts = load_tables(_save(wb, tmp_path))
    assert len(ts) == 2
    assert "Fiscal | 2024" in ts[0].df.columns and "Calendar | 2023" in ts[1].df.columns
    assert ts[1].df["Calendar | 2023"].tolist() == [9]


def test_title_row_flat_header(tmp_path):
    """A caption above a flat header is not a header, and the header is not data."""
    wb = openpyxl.Workbook(); ws = wb.active
    ws["A1"] = "Bảng 3. Doanh thu theo năm"
    ws["A2"] = "Chỉ tiêu"; ws["B2"] = "Tổng"; ws["C2"] = 2019; ws["D2"] = 2020
    ws["A3"] = "Bán lẻ"; ws["B3"] = 100; ws["C3"] = 10; ws["D3"] = 20
    ws["A4"] = "Bán sỉ"; ws["B4"] = 40; ws["C4"] = 5; ws["D4"] = 7
    t = load_tables(_save(wb, tmp_path, "a.xlsx"))[0]
    assert "Bảng 3. Doanh thu theo năm" in t.title
    assert _cols(t) == ["Chỉ tiêu", "Tổng", "2019", "2020"]
    assert t.df["Chỉ tiêu"].tolist() == ["Bán lẻ", "Bán sỉ"]
    assert t.df["2020"].tolist() == [20, 7]
    q = "What is the sum of column 2020? Answer with only the integer."
    assert _ask([t], q)[0] == "27"
    q = "Which two rows have the highest 2020? Answer with the names separated by a comma and one space."
    assert _ask([t], q)[0] == "Bán lẻ, Bán sỉ"

    # Year banner with no corner label: every year stays a value column.
    wb = openpyxl.Workbook(); ws = wb.active
    ws["A1"] = "Bảng 3. Doanh thu theo năm"
    ws["A2"] = 2019; ws["B2"] = 2020; ws["C2"] = 2021
    ws["A3"] = 10; ws["B3"] = 20; ws["C3"] = 30
    ws["A4"] = 5; ws["B4"] = 7; ws["C4"] = 12
    t = load_tables(_save(wb, tmp_path, "a2.xlsx"))[0]
    assert "Bảng 3" in t.title
    assert _values(t) == ["2019", "2020", "2021"]
    assert len(t.df) == 2
    assert t.df["2019"].tolist() == [10, 5]
    assert t.df["2020"].tolist() == [20, 7]
    assert _ask([t], "What is the sum of column 2020? Answer with only the integer.")[0] == "27"

    # A one-cell caption that is not in the label column is still a caption.
    wb = openpyxl.Workbook(); ws = wb.active
    ws["B1"] = "Bảng 3. Doanh thu theo năm"
    ws["A2"] = "Chỉ tiêu"; ws["B2"] = 2019; ws["C2"] = 2020; ws["D2"] = 2021
    ws["A3"] = "Bán lẻ"; ws["B3"] = 10; ws["C3"] = 20; ws["D3"] = 30
    ws["A4"] = "Bán sỉ"; ws["B4"] = 5; ws["C4"] = 7; ws["D4"] = 12
    t = load_tables(_save(wb, tmp_path, "a3.xlsx"))[0]
    assert "Bảng 3. Doanh thu theo năm" in t.title
    assert _cols(t) == ["Chỉ tiêu", "2019", "2020", "2021"]
    assert _ask([t], "What is the sum of column 2019? Answer with only the integer.")[0] == "15"


def test_title_row_two_level_header(tmp_path):
    """Title plus a parent row plus a year row. Years stay under the parent, not in the data."""
    wb = openpyxl.Workbook(); ws = wb.active
    ws["A1"] = "Bảng 3. Doanh thu theo năm"
    ws["A2"] = "Chỉ tiêu"
    ws["B2"] = "Doanh thu"; ws.merge_cells("B2:C2")
    ws["D2"] = "Chi phí"; ws.merge_cells("D2:E2")
    ws["A3"] = "Năm"; ws["B3"] = 2019; ws["C3"] = 2020; ws["D3"] = 2019; ws["E3"] = 2020
    ws["A4"] = "Bán lẻ"; ws["B4"] = 10; ws["C4"] = 20; ws["D4"] = 3; ws["E4"] = 4
    ws["A5"] = "Bán sỉ"; ws["B5"] = 5; ws["C5"] = 7; ws["D5"] = 1; ws["E5"] = 2
    t = load_tables(_save(wb, tmp_path))[0]
    assert "Bảng 3" in t.title
    assert _cols(t) == ["Chỉ tiêu", "Doanh thu | 2019", "Doanh thu | 2020", "Chi phí | 2019", "Chi phí | 2020"]
    assert t.df["Chỉ tiêu"].tolist() == ["Bán lẻ", "Bán sỉ"]
    assert t.df["Doanh thu | 2020"].tolist() == [20, 7]
    got, pl = _ask([t], "What is the sum of Doanh thu 2020? Answer with only the integer.")
    assert pl["columns"] == ["Doanh thu | 2020"] and got == "27"
    got, pl = _ask([t], "What is the sum of Doanh thu? Answer with only the integer.")
    assert pl["columns"] == ["Doanh thu | *"] and got == "42"
    got, pl = _ask([t], "Which two rows have the highest Doanh thu 2020? Answer with the names separated by a comma and one space.")
    assert got == "Bán lẻ, Bán sỉ"


def test_two_title_rows_blank_and_trailing_note(tmp_path):
    """Two captions, a blank row, then the header. A trailing source line is a note."""
    wb = openpyxl.Workbook(); ws = wb.active
    ws["A1"] = "Bảng 3. Doanh thu theo năm"
    ws["D2"] = "Đơn vị tính: tỷ đồng"
    ws["A4"] = "Chỉ tiêu"; ws["B4"] = 2019; ws["C4"] = 2020; ws["D4"] = 2021
    ws["A5"] = "Bán lẻ"; ws["B5"] = 10; ws["C5"] = 20; ws["D5"] = 30
    ws["A6"] = "Bán sỉ"; ws["B6"] = 5; ws["C6"] = 7; ws["D6"] = 12
    ws["A8"] = "Nguồn: Tổng cục Thống kê"
    t = load_tables(_save(wb, tmp_path))[0]
    assert "Bảng 3. Doanh thu theo năm" in t.title and "Đơn vị tính: tỷ đồng" in t.title
    assert _cols(t) == ["Chỉ tiêu", "2019", "2020", "2021"]
    assert t.df["Chỉ tiêu"].tolist() == ["Bán lẻ", "Bán sỉ"]
    assert any("Nguồn" in n for n in t.notes)
    assert _ask([t], "What is the sum of column 2021? Answer with only the integer.")[0] == "42"
    assert _ask([t], "Which two rows have the highest 2020? Answer with the names separated by a comma and one space.")[0] == "Bán lẻ, Bán sỉ"


def test_flat_header_without_title(tmp_path):
    """A flat one-row header is unchanged when nothing sits above it."""
    wb = openpyxl.Workbook(); ws = wb.active
    ws["A1"] = "Name"; ws["B1"] = "Val"
    ws["A2"] = "Total"; ws["B2"] = 3
    for n, v in (("  A", 1), ("  B", 2)):
        ws.append([n, v])
    t = load_tables(_save(wb, tmp_path))[0]
    assert t.title == ""
    assert t.header_rows == [1]
    assert _cols(t) == ["Name", "Val"]
    assert t.df["_row_path"].tolist() == ["Total", "Total > A", "Total > B"]
    assert t.df["Val"].tolist() == [3, 1, 2]
    assert _ask([t], "What is the sum of Val? Answer with only the integer.")[0] == "3"
    assert _ask([t], "Which name has the highest Val? Answer with only the name.")[0] == "B"


def test_multilevel_header_without_title(tmp_path):
    """Parent | year headers, including the wildcard over the children, with no caption above."""
    wb = openpyxl.Workbook(); ws = wb.active
    ws["A1"] = "Item"; ws.merge_cells("A1:A2")
    ws["B1"] = "Revenue"; ws.merge_cells("B1:C1")
    ws["B2"] = 2019; ws["C2"] = 2020
    ws["A3"] = "Sales"; ws["B3"] = 10; ws["C3"] = 20
    ws["A4"] = "Services"; ws["B4"] = 5; ws["C4"] = 7
    t = load_tables(_save(wb, tmp_path))[0]
    assert t.title == ""
    assert t.header_rows == [1, 2]
    assert _cols(t) == ["Item", "Revenue | 2019", "Revenue | 2020"]
    assert t.df["Item"].tolist() == ["Sales", "Services"]
    got, pl = _ask([t], "What is the sum of Revenue 2020? Answer with only the integer.")
    assert pl["columns"] == ["Revenue | 2020"] and got == "27"
    got, pl = _ask([t], "What is the sum of Revenue? Answer with only the integer.")
    assert pl["columns"] == ["Revenue | *"] and got == "42"
    assert _ask([t], "Which item has the highest Revenue 2020? Answer with only the name.")[0] == "Sales"


def test_record_table_with_dates(tmp_path):
    import datetime as dt
    wb = openpyxl.Workbook(); ws = wb.active
    ws["A1"] = "No"; ws.merge_cells("A1:A2"); ws["B1"] = "Name"; ws.merge_cells("B1:C1")
    ws["B2"] = "First"; ws["C2"] = "Last"; ws["D1"] = "Born"; ws.merge_cells("D1:D2")
    ws.append([1, "An", "Le", dt.datetime(2000, 1, 1)])
    ws.append([2, "Binh", "Tran", dt.datetime(1999, 5, 2)])
    t = load_tables(_save(wb, tmp_path))[0]
    assert len(t.df) == 2
    assert "Name | Last" in t.df.columns and "Born" in t.df.columns
    assert t.df["Name | Last"].tolist() == ["Le", "Tran"]
