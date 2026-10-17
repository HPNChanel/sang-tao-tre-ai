"""Independent arithmetic, cell provenance and bounded transport counterexamples."""
import datetime as dt
from decimal import Decimal
import json
from pathlib import Path
import sys

import openpyxl
import pytest

import run
from src import llm_client as L
from src.cellplan_v2.workbook import WorkbookIndex, percent_format
from src.cellplan_v2.context import build_context
from src.cellplan_v2.execute import execute, expression, PlanError, parse_date
from src.cellplan_v2 import planner


@pytest.fixture(autouse=True)
def offline(monkeypatch):
    monkeypatch.setenv('OPENROUTER_API_KEY','fake-test-key')
    monkeypatch.setattr(L,'MODEL',L.APPROVED_MODEL)
    monkeypatch.setattr(L,'BILLING_HALT_REASON',None)
    monkeypatch.setattr(L,'_post',lambda *a:pytest.fail('unmocked paid request'))


@pytest.fixture
def book(tmp_path):
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'Ledger'
    ws.append(['Person','Team','Value','Birth','Event'])
    ws.append(['Same','North',Decimal('0.1'),dt.date(2000,6,15),dt.date(2020,6,14)])
    ws.append(['Same','South',Decimal('0.2'),dt.date(2000,6,15),dt.date(2020,6,15)])
    ws.append(['Total','All','=SUM(C2:C3)',None,None])
    ws['G2']='=SUM(C2,H2)'; ws['G3']='=C3+H3'; ws.merge_cells('I1:J1'); ws['I1']=10
    ws['K2']='2'; ws['K3']='10'
    other=wb.create_sheet('Lookup');other.append(['Person','Team','Factor'])
    other.append(['Same','North',2]);other.append(['Same','South',3])
    path=tmp_path/'test.xlsx';wb.save(path)
    index=WorkbookIndex(path)
    yield index
    index.close()


def plan(*steps, fmt=None):
    return dict(version=2,action='plan',steps=list(steps),result=steps[-1]['id'],format=fmt or {})


def read(region='A2:E3',**extras):
    return dict(id='r',op='read',sheet=0,range=region,columns={'person':'A','team':'B','v':'C','birth':'D','event':'E'},**extras)


def test_raw_coordinates_exact_decimal_and_uncached_formulas(book):
    assert book.value(0,'C2')==Decimal('0.1')
    assert book.value(0,'C4')==Decimal('0.3')
    assert book.value(0,'G2')==Decimal('0.1')
    assert book.value(0,'G3')==Decimal('0.2')
    p=plan(read(),dict(id='a',op='aggregate',input='r',metrics={'s':{'fn':'sum','column':'v'}}))
    assert execute(p,book)[0]=='0.3'


def test_merged_numeric_exists_once(book):
    p=plan(dict(id='r',op='read',sheet=0,range='I1:J1',columns={'a':'I','b':'J'}),
           dict(id='u',op='unpivot',input='r',columns={'a':'A','b':'B'},key='k',value='v'),
           dict(id='a',op='aggregate',input='u',metrics={'sum':{'fn':'sum','column':'v'}}))
    assert execute(p,book)[0]=='10'
    assert book.value(0,'J1',merged=True) is None


def test_age_is_at_event_and_accounts_for_birthday(book):
    p=plan(read(),dict(id='d',op='derive',input='r',columns={'age':{'op':'years','args':[{'col':'birth'},{'col':'event'}]}}),
           fmt={'columns':['age']})
    assert execute(p,book)[0]=='19, 20'


def test_composite_key_join_and_weighted_sum(book):
    p=plan(read(),dict(id='lookup',op='read',sheet='Lookup',range='A2:C3',columns={'who':'A','unit':'B','factor':'C'}),
           dict(id='j',op='join',inputs=['r','lookup'],on={'person':'who','team':'unit'}),
           dict(id='d',op='derive',input='j',columns={'weighted':{'op':'mul','args':[{'col':'v'},{'col':'r_factor'}]}}),
           dict(id='a',op='aggregate',input='d',metrics={'sum':{'fn':'sum','column':'weighted'}}))
    assert execute(p,book)[0]=='0.8'
    p['steps'][2]['on']={'person':'who'}
    with pytest.raises(PlanError,match='not unique'): execute(p,book)


def test_numeric_text_ranking_requires_declared_type(book):
    p=plan(dict(id='r',op='read',sheet=0,range='K2:K3',columns={'v':'K'}),
           dict(id='s',op='sort',input='r',by=[{'column':'v','descending':True}]),
           dict(id='t',op='take',input='s'),fmt={'columns':['v']})
    with pytest.raises(PlanError,match='explicit sort type'): execute(p,book)
    p['steps'][1]['by'][0]['type']='number'
    assert execute(p,book)[0]=='10'


def test_percent_points_and_relative_change_are_distinct():
    row={'a':Decimal('.20'),'b':Decimal('.25')}
    args=[{'col':'a'},{'col':'b'}]
    assert expression({'op':'change_ratio','args':args},row)==Decimal('.25')
    assert expression({'op':'sub','args':list(reversed(args))},row)==Decimal('.05')


def test_conditional_unit_conversion_does_not_evaluate_unused_branch():
    e={'op':'if','where':{'column':'unit','cmp':'eq','value':'mL'},
       'then':{'op':'div','args':[{'col':'v'},{'literal':1000}]},'else':{'col':'v'}}
    assert expression(e,{'unit':'mL','v':Decimal('1250')})==Decimal('1.25')
    assert expression(e,{'unit':'L','v':Decimal('1.25')})==Decimal('1.25')


def test_locale_dates_are_explicit():
    assert parse_date('03/04/2026','dmy')==dt.date(2026,4,3)
    assert parse_date('03/04/2026','mdy')==dt.date(2026,3,4)
    with pytest.raises(PlanError): parse_date('03/04/2026')


def test_complete_context_keeps_cell_and_formula_metadata(book):
    context=build_context(book,'Value?',dual=True)
    assert context['complete'] and context['omitted_rows']==0
    rows={r['r']:r for r in context['views'][0]['rows']}
    assert rows[2]['C']=={'n':'0.1'}
    assert rows[4]['C']['v']=={'n':'0.3'} and rows[4]['C']['f']=='SUM(C2:C3)'
    assert context['structure_hints_untrusted']


@pytest.fixture
def marks(tmp_path):
    """A statistical table that marks unavailable figures with '..' and 'x'."""
    wb=openpyxl.Workbook();ws=wb.active
    for row in [('Region','Y2019','Y2020','Grade'),('North',10,'..','M'),('South','x',4,'F'),('East',6,8,'F')]:
        ws.append(row)
    path=tmp_path/'marks.xlsx';wb.save(path)
    index=WorkbookIndex(path)
    yield index
    index.close()


def stats(*steps, fmt=None):
    return plan(dict(id='r',op='read',sheet=0,range='A2:D4',columns={'region':'A','a':'B','b':'C','grade':'D'}),
                *steps,fmt=fmt)


def test_missing_marks_propagate_through_arithmetic_and_aggregates_skip_them(marks):
    assert expression({'op':'add','args':[{'literal':'..'},{'literal':1}]},{}) is None
    assert execute(stats(dict(id='s',op='aggregate',input='r',metrics={'m':{'fn':'mean','column':'b'}})),marks)[0]=='6'
    change=stats(dict(id='d',op='derive',input='r',columns={'c':{'op':'sub','args':[{'col':'b'},{'col':'a'}]}}),
                 dict(id='s',op='aggregate',input='d',metrics={'n':{'fn':'sum','column':'c'},'k':{'fn':'count','column':'c'}}),
                 fmt={'columns':['n','k']})
    assert execute(change,marks)[0]=='2, 1'


@pytest.mark.parametrize('descending,order',[(True,'East, South, North'),(False,'South, East, North')])
def test_missing_marks_sort_last_in_either_direction(marks,descending,order):
    p=stats(dict(id='s',op='sort',input='r',by=[{'column':'b','descending':descending}]),fmt={'columns':['region']})
    assert execute(p,marks)[0]==order


def test_text_sort_keeps_values_that_look_like_marks(marks):
    p=stats(dict(id='s',op='sort',input='r',by=[{'column':'grade'},{'column':'region'}]),fmt={'columns':['region']})
    assert execute(p,marks)[0]=='East, South, North'


def test_sort_on_a_column_without_values_is_an_error(marks):
    p=plan(dict(id='r',op='read',sheet=0,range='A2:B3',columns={'region':'A','a':'B'}),
           dict(id='f',op='filter',input='r',where={'column':'region','cmp':'eq','value':'South'}),
           dict(id='s',op='sort',input='f',by=[{'column':'a','type':'number'}]))
    with pytest.raises(PlanError,match="'a' has no values"): execute(p,marks)


def test_missing_marks_never_satisfy_numeric_comparisons(marks):
    p=stats(dict(id='f',op='filter',input='r',where={'column':'b','cmp':'gt','value':5}),fmt={'columns':['region']})
    assert execute(p,marks)[0]=='East'


def test_non_numeric_operands_and_unknown_ops_are_named():
    with pytest.raises(PlanError,match="numeric operand unavailable: 'abc'"):
        expression({'op':'add','args':[{'literal':'abc'},{'literal':1}]},{})
    with pytest.raises(PlanError,match="unsupported expression op 'pow'"):
        expression({'op':'pow','args':[{'literal':2},{'literal':3}]},{})


def test_format_scale_is_rejected_with_derive_guidance(book):
    p=plan(read(),dict(id='a',op='aggregate',input='r',metrics={'s':{'fn':'sum','column':'v'}}),fmt={'scale':100})
    with pytest.raises(PlanError,match="unsupported format fields \\['scale'\\].*derive step"): execute(p,book)


def test_rejections_name_the_offending_fields_and_the_allowed_form(book):
    where={'column':'team','cmp':'eq','value':'North'}
    p=plan(read(),dict(id='a',op='aggregate',input='r',where=where,metrics={'s':{'fn':'sum','column':'v'}}))
    with pytest.raises(PlanError,match='unexpected fields for aggregate: where; aggregate takes id, op, '
                                       'group_by, input, metrics; filter the input first'): execute(p,book)
    p=plan(read(),dict(id='f',op='filter',input='r',where=dict(where,col='v')))
    with pytest.raises(PlanError,match='unknown filter field col; a predicate takes column'): execute(p,book)
    p=plan(read(),dict(id='d',op='derive',input='r',columns={'x':'v*2'}))
    with pytest.raises(PlanError,match="expression must be an object, not 'v\\*2'"): execute(p,book)
    p=plan(read(),dict(id='d',op='derive',input='r',columns={'x':{'column':'v'}}))
    with pytest.raises(PlanError,match='invalid expression fields column; an expression is'): execute(p,book)
    p=plan(read(),dict(id='a',op='aggregate',input='r',metrics={'s':{'fn':'sum','col':'v'}}))
    with pytest.raises(PlanError,match="invalid metric fields in 's'; a metric takes fn, column"): execute(p,book)


@pytest.fixture
def side_tables(tmp_path):
    """Receipts in A:D, a titled price catalog in K:M of the same sheet (rows do not correspond)."""
    wb=openpyxl.Workbook();ws=wb.active
    ws['A1']='receipts';ws['K2']='catalog'
    rows=[('Date','Site','Code','Qty'),('2025-01-05','North','P1',2),('2025-02-07','North','P2',3),
          ('2025-02-09','South','P1',4),('2025-03-01','North','P1',1)]
    for i,row in enumerate(rows,start=2):
        for col,value in zip('ABCD',row): ws[f'{col}{i}']=value
    for i,row in enumerate([('Kind','Price','Code'),('Wood',5,'P1'),('Stone',7,'P2')],start=3):
        for col,value in zip('KLM',row): ws[f'{col}{i}']=value
    path=tmp_path/'side.xlsx';wb.save(path)
    index=WorkbookIndex(path)
    yield index
    index.close()


def test_side_by_side_tables_are_separate_geometric_blocks(side_tables,book):
    from src.cellplan_v2.context import table_blocks
    assert [b['range'] for b in table_blocks(side_tables)]==['A1:D6','K2:M5']
    assert build_context(side_tables,'Total?')['table_blocks_untrusted'][1]=={'sheet':0,'range':'K2:M5','cells':10}
    assert 'table_blocks_untrusted' not in build_context(book,'Value?') or len(table_blocks(book))>1


def test_side_by_side_lookup_ignores_blank_title_keys(side_tables):
    p=plan(dict(id='r',op='read',sheet=0,range='A3:D6',columns={'date':'A','site':'B','code':'C','qty':'D'}),
           dict(id='c',op='read',sheet=0,range='K2:M5',columns={'kind':'K','price':'L','code':'M'}),
           dict(id='j',op='join',inputs=['r','c'],on={'code':'code'}),
           dict(id='f',op='filter',input='j',where={'column':'site','cmp':'eq','value':'North'}),
           dict(id='d',op='derive',input='f',columns={'value':{'op':'mul','args':[{'col':'qty'},{'col':'r_price'}]}}),
           dict(id='a',op='aggregate',input='d',metrics={'total':{'fn':'sum','column':'value'}}))
    assert execute(p,side_tables)[0]=='36'
    # Column letters outside the range's own columns still read the requested rows.
    p['steps'][1]['range']='K2:K5'
    assert execute(p,side_tables)[0]=='36'


def test_conditional_metrics_and_cross_join_combine_periods(side_tables):
    month=lambda m:{'column':'month','cmp':'eq','value':m}
    p=plan(dict(id='r',op='read',sheet=0,range='A3:D6',columns={'date':'A','site':'B','code':'C','qty':'D'}),
           dict(id='m',op='derive',input='r',columns={'month':{'op':'month','args':[{'col':'date'}]}}),
           dict(id='a',op='aggregate',input='m',metrics={'jan':{'fn':'sum','column':'qty','where':month('2025-01')},
                                                         'feb':{'fn':'sum','column':'qty','where':month('2025-02')}}),
           dict(id='d',op='derive',input='a',columns={'change':{'op':'sub','args':[{'col':'feb'},{'col':'jan'}]}}),
           fmt={'columns':['change']})
    assert execute(p,side_tables)[0]=='5'
    q=plan(dict(id='r',op='read',sheet=0,range='A3:D6',columns={'site':'B','qty':'D'}),
           dict(id='n',op='filter',input='r',where={'column':'site','cmp':'eq','value':'North'}),
           dict(id='s',op='filter',input='r',where={'column':'site','cmp':'eq','value':'South'}),
           dict(id='na',op='aggregate',input='n',metrics={'north':{'fn':'sum','column':'qty'}}),
           dict(id='sa',op='aggregate',input='s',metrics={'south':{'fn':'sum','column':'qty'}}),
           dict(id='x',op='join',inputs=['na','sa'],how='cross'),
           dict(id='d',op='derive',input='x',columns={'gap':{'op':'sub','args':[{'col':'north'},{'col':'r_south'}]}}),
           fmt={'columns':['gap']})
    assert execute(q,side_tables)[0]=='2'


def test_derive_combines_single_row_steps(side_tables):
    p=plan(dict(id='r',op='read',sheet=0,range='A3:D6',columns={'site':'B','qty':'D'}),
           dict(id='n',op='filter',input='r',where={'column':'site','cmp':'eq','value':'North'}),
           dict(id='s',op='filter',input='r',where={'column':'site','cmp':'eq','value':'South'}),
           dict(id='na',op='aggregate',input='n',metrics={'north':{'fn':'sum','column':'qty'}}),
           dict(id='sa',op='aggregate',input='s',metrics={'south':{'fn':'sum','column':'qty'}}),
           dict(id='d',op='derive',inputs=['na','sa'],columns={'gap':{'op':'sub','args':[{'col':'north'},{'col':'south'}]}}),
           fmt={'columns':['gap']})
    assert execute(p,side_tables)[0]=='2'
    p['steps'][-1]['inputs']=['na','n']
    with pytest.raises(PlanError,match="step 'n' has 3 rows"): execute(p,side_tables)
    p['steps'][-1]['inputs']=['na','na']
    with pytest.raises(PlanError,match="share column 'north'"): execute(p,side_tables)
    p['steps'][-1].update(inputs=['na','sa'],input='na')
    with pytest.raises(PlanError,match='not both'): execute(p,side_tables)


def test_weighted_mean_questions_need_value_times_weight():
    ratio=plan(dict(id='r',op='read',sheet=0,range='A2:B9',columns={'v':'A','w':'B'}),
               dict(id='a',op='aggregate',input='r',metrics={'s':{'fn':'sum','column':'v'},'t':{'fn':'sum','column':'w'}}),
               dict(id='d',op='derive',input='a',columns={'m':{'op':'div','args':[{'col':'s'},{'col':'t'}]}}))
    for query in ('Điểm trung bình có trọng số là bao nhiêu?','What is the weighted average score?','가중 평균은?'):
        with pytest.raises(PlanError,match='weighted mean'): planner.check_weighted(query,ratio)
    planner.check_weighted('What is the mean of the ratio?',ratio)
    ratio['steps'][1]['metrics']={'m':{'fn':'weighted_mean','column':'v','weight':'w'}}
    planner.check_weighted('What is the weighted average score?',ratio)


def test_weighted_mean_question_repairs_a_ratio_of_sums(side_tables,monkeypatch):
    steps=[dict(id='r',op='read',sheet=0,range='K4:M5',columns={'price':'L','code':'M'}),
           dict(id='w',op='read',sheet=0,range='C3:D6',columns={'code':'C','qty':'D'}),
           dict(id='j',op='join',inputs=['w','r'],on={'code':'code'})]
    ratio=plan(*steps,dict(id='a',op='aggregate',input='j',metrics={'p':{'fn':'sum','column':'r_price'},'q':{'fn':'sum','column':'qty'}}),
               dict(id='d',op='derive',input='a',columns={'m':{'op':'div','args':[{'col':'p'},{'col':'q'}]}}),fmt={'columns':['m']})
    good=plan(*steps,dict(id='a',op='aggregate',input='j',metrics={'m':{'fn':'weighted_mean','column':'r_price','weight':'qty'}}))
    replies=iter([ratio,good]);feedback=[]
    def post(body,timeout):
        sent=json.loads(body)['messages'];feedback.append(sent[-1]['content'] if len(sent)>2 else '')
        return response(next(replies))
    monkeypatch.setattr(L,'_post',post)
    cache=run.TableCache();b=L.Budget()
    try: assert planner.answer('What is the weighted average price by quantity? Return two decimal places.',
                               side_tables.path,cache,b,'raw','q')=='5.60'
    finally: planner.close_cache(cache)
    assert b.attempts==2 and 'weighted mean' in feedback[1]


@pytest.mark.parametrize('text,query,expected',[
    ('12.5%','What share of the units are wood?','12.5'),
    ('-3%','Tỷ lệ thay đổi là bao nhiêu phần trăm?','-3'),
    ('12.5%','Trả lời kèm ký hiệu %.','12.5%'),
    ('12.5%, 13%','List the shares.','12.5%, 13%'),
    ('North','Which site?','North'),
])
def test_a_lone_percentage_drops_an_unrequested_sign(text,query,expected):
    assert planner.drop_unasked_percent(text,query)==expected


def test_weighted_mean_uses_value_times_weight(side_tables):
    p=plan(dict(id='r',op='read',sheet=0,range='K4:M5',columns={'price':'L','code':'M'}),
           dict(id='w',op='read',sheet=0,range='C3:D6',columns={'code':'C','qty':'D'}),
           dict(id='j',op='join',inputs=['w','r'],on={'code':'code'}),
           dict(id='a',op='aggregate',input='j',metrics={'m':{'fn':'weighted_mean','column':'r_price','weight':'qty'}}),
           fmt={'decimals':2})
    assert execute(p,side_tables)[0]=='5.60'
    for metric in ({'fn':'mean','column':'r_price','weight':'qty'},{'fn':'sum','weight':'qty'},
                   {'fn':'weighted_mean','weight':'qty'},{'fn':'weighted_mean','column':'r_price'}):
        p['steps'][3]['metrics']['m']=metric
        with pytest.raises(PlanError,match='weighted_mean'): execute(p,side_tables)


def test_join_keeps_plain_right_aliases_without_collisions(side_tables):
    # L (price) is deliberately aliased as site: the left site keeps its name, the right one is r_site.
    p=plan(dict(id='r',op='read',sheet=0,range='A3:D6',columns={'site':'B','code':'C','qty':'D'}),
           dict(id='c',op='read',sheet=0,range='K2:M5',columns={'kind':'K','site':'L','code':'M'}),
           dict(id='j',op='join',inputs=['r','c'],on={'code':'code'}),
           dict(id='f',op='filter',input='j',where={'all':[{'column':'kind','cmp':'eq','value':'Wood'},
                                                           {'column':'site','cmp':'eq','value':'North'}]}),
           dict(id='d',op='derive',input='f',columns={'value':{'op':'mul','args':[{'col':'qty'},{'col':'r_site'}]}}),
           dict(id='a',op='aggregate',input='d',metrics={'total':{'fn':'sum','column':'value'}}))
    assert execute(p,side_tables)[0]=='15'


def test_malformed_plan_errors_get_repair_feedback(book,monkeypatch):
    replies=iter([plan(dict(id='c',op='cell',sheet=0,cell='C4'))]*2)
    real,calls=planner.execute,[]
    def flaky(p,b,*rest):
        calls.append(p)
        if len(calls)==1: raise AttributeError("'list' object has no attribute 'get'")
        return real(p,b,*rest)
    monkeypatch.setattr(planner,'execute',flaky)
    monkeypatch.setattr(L,'_post',lambda body,timeout:response(next(replies)))
    cache=run.TableCache();b=L.Budget()
    try: assert planner.answer('What is the total?',book.path,cache,b,'raw','q')=='0.3'
    finally: planner.close_cache(cache)
    assert b.attempts==2 and b.diagnostics['status']=='executed'


def test_unknown_columns_name_the_available_aliases(side_tables):
    p=plan(dict(id='r',op='read',sheet=0,range='A3:D6',columns={'site':'B','qty':'D'}),
           dict(id='d',op='derive',input='r',columns={'x':{'op':'mul','args':[{'col':'amount'},{'literal':2}]}}))
    with pytest.raises(PlanError,match="'amount'.*qty, site"): execute(p,side_tables)


@pytest.mark.parametrize('query,places',[
    ('Dựa trên sổ tính, tổng là bao nhiêu? Chỉ ghi số nguyên.',0),
    ('Dựa trên sổ tính, trung bình là bao nhiêu? Ghi đúng 2 chữ số thập phân.',2),
    ('Using only the workbook, what is the total? Exclude absences; return an integer.',0),
    ('What is the mean? Return two decimal places.',2),
    ('Which sites? Answer with the names separated by a comma and one space.',None),
    ('Liệt kê các số nguyên, cách nhau bởi dấu phẩy.',None),
    ('Who is youngest? Answer with only the last name.',None),
    ('Quelle est la moyenne ? Arrondissez à deux décimales.',2),
    ('Donnez la réponse sous forme de nombre entier.',0),
    ('평균은 얼마입니까? 소수점 이하 둘째 자리까지 반올림하세요.',2),
    ('평균을 정수로 답하세요.',0),
    ('平均值是多少？保留两位小数。',2),
    ('总数是多少？请给出整数。',0),
    ('Tỷ lệ là bao nhiêu? Làm tròn đến 1 chữ số thập phân.',1),
    ('Quels pays ? Séparez les noms par une virgule.',None),
])
def test_requested_single_number_is_read_from_the_question(query,places):
    assert planner.requested_number(query)==places


@pytest.mark.parametrize('query,keep',[
    ('Tỷ lệ là bao nhiêu phần trăm? Ghi 1 chữ số thập phân, không kèm ký hiệu %.',False),
    ('Give the share with one decimal place, without the percent sign.',False),
    ('What share? Round to one decimal place.',False),
    ('Trả lời kèm ký hiệu %, 1 chữ số thập phân.',True),
    ('Donnez le taux avec une décimale, avec le signe %.',True),
])
def test_percent_sign_is_printed_only_when_asked(query,keep):
    assert planner.keeps_percent_sign(query) is keep
    assert planner.single_number('81.25%',1,planner.keeps_percent_sign(query))==('81.3%' if keep else '81.3')


def test_single_number_questions_reject_lists_and_apply_precision(side_tables,monkeypatch):
    rows=plan(dict(id='r',op='read',sheet=0,range='A3:D6',columns={'qty':'D'}),fmt={'columns':['qty']})
    total=plan(dict(id='r',op='read',sheet=0,range='A3:D6',columns={'qty':'D'}),
               dict(id='a',op='aggregate',input='r',metrics={'m':{'fn':'mean','column':'qty'}}))
    replies=iter([rows,total]);feedback=[]
    def post(body,timeout):
        sent=json.loads(body)['messages'];feedback.append(sent[-1]['content'] if len(sent)>2 else '')
        return response(next(replies))
    monkeypatch.setattr(L,'_post',post)
    cache=run.TableCache();b=L.Budget()
    try: assert planner.answer('What is the mean quantity? Return two decimal places.',side_tables.path,cache,b,'raw','q')=='2.50'
    finally: planner.close_cache(cache)
    assert b.attempts==2 and 'has 4 rows' in feedback[1]


def test_clock_time_cells_render_as_json_context(tmp_path):
    wb=openpyxl.Workbook();ws=wb.active
    ws.append(['Shift','Start']);ws.append(['Early',dt.time(8,30)])
    path=tmp_path/'times.xlsx';wb.save(path)
    index=WorkbookIndex(path)
    try:
        context=build_context(index,'When does Early start?')
        assert context['complete'] and context['views'][0]['rows'][1]['B']['v']=={'time':'08:30:00'}
    finally: index.close()


def test_internal_v2_error_is_recorded_not_raised(book,monkeypatch):
    def broken(*a,**k): raise TypeError('unexpected cell type')
    monkeypatch.setattr(planner,'build_context',broken)
    cache=run.TableCache();b=L.Budget()
    try: assert planner.answer('Total?',book.path,cache,b,'raw','q') is None
    finally: planner.close_cache(cache)
    assert b.diagnostics['status']=='index_or_context_error' and b.diagnostics['error'].startswith('TypeError')


def test_unpivot_limit_checked_before_allocating(book,monkeypatch):
    import src.cellplan_v2.execute as engine
    monkeypatch.setattr(engine,'MAX_ROWS',3)
    p=plan(dict(id='r',op='read',sheet=0,range='A2:C3',columns={'x':'B','y':'C'}),
           dict(id='u',op='unpivot',input='r',columns={'x':'X','y':'Y'}))
    with pytest.raises(PlanError,match='unpivot record limit'): execute(p,book)


def response(obj,cost=.0001):
    return {'choices':[{'finish_reason':'stop','message':{'content':json.dumps(obj)}}],
            'usage':{'cost':cost,'prompt_tokens':100,'completion_tokens':100}}


def test_inspect_repair_and_execution_share_three_attempts(book,monkeypatch):
    good=plan(dict(id='c',op='cell',sheet=0,cell='C4'))
    replies=iter([{'version':2,'action':'inspect','regions':[{'sheet':0,'range':'A1:C4'}]},
                  plan(dict(id='c',op='cell',sheet=0,cell='ZZ999')),good])
    bodies=[]
    def post(body,timeout):
        bodies.append(json.loads(body));return response(next(replies))
    monkeypatch.setattr(L,'_post',post)
    cache=run.TableCache();budget=L.Budget()
    try:
        assert planner.answer('What is the total?',book.path,cache,budget,'dual','q')=='0.3'
    finally: planner.close_cache(cache)
    assert budget.attempts==3 and budget.actual_usd==pytest.approx(.0003)
    assert all(len(json.dumps(b,ensure_ascii=False).encode())<=L.MAX_BODY_BYTES for b in bodies)
    assert all(b['messages'][1]['content'].endswith('\n\nQUESTION:\nWhat is the total?') for b in bodies)


def test_an_abstention_is_asked_again_once_with_the_full_workbook_note(book,monkeypatch):
    replies=iter([{'version':2,'action':'abstain'},plan(dict(id='c',op='cell',sheet=0,cell='C4'))])
    bodies=[]
    def post(body,timeout):
        bodies.append(json.loads(body));return response(next(replies))
    monkeypatch.setattr(L,'_post',post)
    cache=run.TableCache();budget=L.Budget()
    try: assert planner.answer('What is the total?',book.path,cache,budget,'raw','q')=='0.3'
    finally: planner.close_cache(cache)
    assert budget.attempts==2 and budget.diagnostics['status']=='executed'
    assert budget.diagnostics['reasked_abstention'] is True
    assert [e['status'] for e in budget.log]==['abstained','executed']
    assert len(bodies[0]['messages'])==2 and planner.REASK_ABSTENTION in bodies[1]['messages'][2]['content']


@pytest.mark.parametrize('profile,attempts',[('raw',2),('direct',1)])
def test_a_repeated_abstention_ends_v2_and_leaves_the_rest_of_the_budget(book,monkeypatch,profile,attempts):
    monkeypatch.setattr(L,'_post',lambda body,timeout:response({'version':2,'action':'abstain'}))
    cache=run.TableCache();budget=L.Budget()
    try: assert planner.answer('What is the total?',book.path,cache,budget,profile,'q') is None
    finally: planner.close_cache(cache)
    assert budget.attempts==attempts and budget.diagnostics['status']=='model_abstained' and budget.allow(0)


@pytest.fixture
def long_book(tmp_path):
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'Rows'
    ws.append(['Code','Region','Amount','Note'])
    for i in range(400):
        ws.append([f'K{i:04d}',['North','South','East'][i%3],i*7,f'entry number {i} of the long ledger'])
    path=tmp_path/'long.xlsx';wb.save(path)
    return path


def test_an_oversized_repair_request_rebuilds_a_smaller_view_without_an_attempt(long_book,monkeypatch):
    replies=iter([plan(dict(id='c',op='cell',sheet=0,cell='ZZ999')),plan(dict(id='c',op='cell',sheet=0,cell='C12'))])
    bodies=[]
    def post(body,timeout):
        if not bodies:  # the first request fits; its repair with feedback does not
            monkeypatch.setattr(L,'MAX_BODY_BYTES',len(body)+40)
        bodies.append(body);return response(next(replies))
    monkeypatch.setattr(L,'_post',post)
    cache=run.TableCache();budget=L.Budget()
    try: assert planner.answer('What is the amount in row 12?',long_book,cache,budget,'raw','q')=='70'
    finally: planner.close_cache(cache)
    assert budget.attempts==2 and budget.diagnostics['repair_view_bytes']==7800
    assert len(bodies[1])<len(bodies[0])<=L.MAX_BODY_BYTES-40
    repair=json.loads(bodies[1])['messages']
    assert repair[1]['content'].endswith('\n\nQUESTION:\nWhat is the amount in row 12?')
    assert repair[2]['content'].startswith('Validation feedback for your previous attempt:')


def test_a_first_request_over_the_cap_still_ends_v2_unsent(long_book,monkeypatch):
    monkeypatch.setattr(L,'MAX_BODY_BYTES',2000)
    cache=run.TableCache();budget=L.Budget()
    try: assert planner.answer('What is the amount in row 12?',long_book,cache,budget,'raw','q') is None
    finally: planner.close_cache(cache)
    assert budget.attempts==0 and budget.diagnostics['status']=='request_exceeds_24_kib'
    assert 'repair_view_bytes' not in budget.diagnostics


def test_repair_view_shrinking_is_bounded(book,monkeypatch):
    bodies=[]
    def post(body,timeout):
        monkeypatch.setattr(L,'MAX_BODY_BYTES',len(body)+40)
        bodies.append(body);return response(plan(dict(id='c',op='cell',sheet=0,cell='ZZ999')))
    monkeypatch.setattr(L,'_post',post)
    built=[]
    real=planner.build_context
    monkeypatch.setattr(planner,'build_context',lambda *a,**k:(built.append(k['max_bytes']),real(*a,**k))[1])
    cache=run.TableCache();budget=L.Budget()
    try: assert planner.answer('What is the total?',book.path,cache,budget,'raw','q') is None
    finally: planner.close_cache(cache)
    # The small workbook already fits, so a smaller view cannot make room for the feedback.
    assert len(bodies)==1 and budget.attempts==1 and budget.diagnostics['status']=='request_exceeds_24_kib'
    assert built==[13000,7800,4680,3000]


def test_direct_answer_is_eligible_with_valid_evidence(book,monkeypatch):
    reply={'version':2,'action':'answer','answer':'0.3','evidence':[{'sheet':0,'range':'C2:C3'}]}
    monkeypatch.setattr(L,'_post',lambda *a:response(reply))
    cache=run.TableCache();b=L.Budget()
    try: assert planner.answer('Total?',book.path,cache,b,'direct','q')=='0.3'
    finally: planner.close_cache(cache)
    assert b.diagnostics['status']=='direct_answer'


def test_model_and_base_url_are_not_silently_replaced(monkeypatch):
    saved=run._apply_openrouter_cli('fake','google/gemini-2.5-flash-lite','https://example.invalid/v1')
    seen=[]
    monkeypatch.setattr(L,'_post',lambda body,timeout:(seen.append(json.loads(body)) or response({'version':2,'action':'abstain'})))
    try:
        b=L.Budget();assert L.request_json([{'role':'user','content':'Q'}],b)[0]['action']=='abstain'
        assert seen[0]['model']=='google/gemini-2.5-flash-lite'
        assert L.URL=='https://example.invalid/v1/chat/completions'
        assert 'only' not in seen[0]['provider']
    finally: run._restore_openrouter_cli(saved)


@pytest.mark.parametrize('cap',[600,1200])
def test_configured_output_cap_reaches_the_transport(monkeypatch,cap):
    seen=[]
    monkeypatch.setattr(L,'MAX_TOKENS',cap)
    monkeypatch.setattr(L,'_post',lambda body,timeout:(seen.append(json.loads(body)) or response({'version':2,'action':'abstain'})))
    L.request_json([{'role':'user','content':'Q'}],L.Budget())
    assert seen[0]['max_tokens']==cap


def test_v2_runtime_integration_uses_existing_transport(book,monkeypatch):
    monkeypatch.setenv('DS2_SOLVER_PROFILE','raw')
    monkeypatch.setattr(L,'_post',lambda *a:response(plan(dict(id='c',op='cell',sheet=0,cell='C4'))))
    cache=run.TableCache()
    try:
        result,b=run.answer_one({'id':'q','query':'What total?','file_name':book.path.name},book.path.parent,cache,timeout=10,use_fork=False)
        assert result=='0.3' and b.attempts==1
    finally: planner.close_cache(cache)


def test_global_stable_question_cost_includes_older_source_and_resharding(tmp_path):
    sys.path.insert(0,str(Path(__file__).parents[1]/'scripts'))
    import qwen_campaign as C
    ledger=C.Ledger.create(tmp_path/'ledger.json','test')
    config={'policy':'selective','output_tokens':1200,'api_enabled':True}
    # Reconstruct a genuine v1 ledger entry, before stable identities existed.
    import hashlib
    hashed=hashlib.sha256(b'x').hexdigest()
    identity={'source_sha':'a'*40,'input_sha256':'b'*64,'configuration':config,'repetition':1,
              'case_ids_sha256':hashlib.sha256(C.canonical([hashed])).hexdigest()}
    old=hashlib.sha256(C.canonical(identity)).hexdigest()
    ledger.state['invocations'][old]=dict(identity,stage='development',cases=[hashed])
    ledger.save()
    for _ in range(2):
        n=ledger.reserve('development',old,'x','c'*64);ledger.finish(n,{'usage':{'cost':'.03'}})
    # Rebinding the same input maps prior spending; changing batch and source
    # cannot obtain another fresh .10 allowance for the same semantic case.
    ledger.register('development','d'*40,'b'*64,config,['x'],case_keys={'x':'e'*64})
    new=ledger.register('development','f'*40,'1'*64,dict(config,solver_profile='dual'),['y'],case_keys={'y':'e'*64})
    n=ledger.reserve('development',new,'y','c'*64);ledger.finish(n,{'usage':{'cost':'.02'}})
    with pytest.raises(C.CampaignError,match='cumulative'): ledger.reserve('development',new,'y','c'*64)
    assert C.money(ledger.totals()['actual_usd'])==Decimal('.08')


def test_sealed_rerun_repetition_keeps_the_stable_question_cap(tmp_path):
    sys.path.insert(0,str(Path(__file__).parents[1]/'scripts'))
    import qwen_campaign as C
    ledger=C.Ledger.create(tmp_path/'ledger.json','test')
    config={'policy':'selective','output_tokens':1200}
    first=ledger.register('sealed','a'*40,'b'*64,config,['x'],case_keys={'x':'e'*64})
    # Re-running a halted split gets its own invocation, but no fresh .10 allowance per case.
    again=ledger.register('sealed','a'*40,'b'*64,config,['x'],2,case_keys={'x':'e'*64})
    assert again!=first
    for _ in range(3):
        n=ledger.reserve('sealed',first,'x','c'*64);ledger.finish(n,{'usage':{'cost':'.03'}})
    with pytest.raises(C.CampaignError,match='cumulative'): ledger.reserve('sealed',again,'x','c'*64)
    with pytest.raises(C.CampaignError,match='repetition'):
        ledger.register('development','a'*40,'b'*64,config,['x'],2,case_keys={'x':'e'*64})


def test_unresolved_billing_stops_other_questions(tmp_path):
    sys.path.insert(0,str(Path(__file__).parents[1]/'scripts'))
    import qwen_campaign as C
    ledger=C.Ledger.create(tmp_path/'ledger.json','test')
    invocation=ledger.register('development','a'*40,'b'*64,{'policy':'selective','output_tokens':1200},['x','y'],
                               case_keys={'x':'d'*64,'y':'e'*64})
    n=ledger.reserve('development',invocation,'x','c'*64);ledger.finish(n,{})
    with pytest.raises(C.CampaignError,match='halted'): ledger.reserve('development',invocation,'y','c'*64)
    assert ledger.totals()['unresolved_usd']=='0.03'


def test_new_paid_registration_cannot_choose_legacy_identity(tmp_path):
    sys.path.insert(0,str(Path(__file__).parents[1]/'scripts'))
    import qwen_campaign as C
    ledger=C.Ledger.create(tmp_path/'ledger.json','test')
    with pytest.raises(C.CampaignError,match='stable case keys'):
        ledger.register('development','a'*40,'b'*64,{'policy':'selective','output_tokens':1200},['x'])
    assert ledger.state['invocations']=={}


@pytest.fixture
def shares(tmp_path):
    """Percent-unit figures stored as text under a % unit row, amounts as text with thousands
    separators, fractions stored with a percent format, and the same fractions without one."""
    wb=openpyxl.Workbook();ws=wb.active
    ws.append(['need','unmet','partial','men','women','share','bare'])
    ws.append([None,'%','%','dollars','dollars',None,None])
    for row in [('any','21.4','22.4','58,176','37,117',0.25,0.25),('info','20.8','8.0','52,516','30,519',0.5,0.5),
                ('counselling','34.1','15.9','71,104','57,842',0.125,0.125)]:
        ws.append(list(row))
    for r in range(3,6): ws[f'F{r}'].number_format='0.0%'
    path=tmp_path/'shares.xlsx';wb.save(path)
    index=WorkbookIndex(path)
    yield index
    index.close()


@pytest.mark.parametrize('fmt,scaled',[('0%',True),('0.0%',True),('#,##0.00%;[Red]-#,##0.00%',True),
                                       ('0.0"%"',False),('0.0\\%',False),('[$%-409]0.0',False),('General',False),(None,False)])
def test_only_an_unquoted_percent_scales_a_number_format(fmt,scaled):
    assert percent_format(fmt) is scaled


def test_column_scales_come_from_formats_text_and_magnitude(shares):
    assert shares.scales(0,3,5,[2,4,6,7,8])=={2:('quantity',0),4:('quantity',0),6:('share',0),7:('small',0),8:None}


def test_figures_already_in_percent_units_are_not_multiplied_again(shares):
    p=plan(dict(id='r',op='read',sheet=0,range='A3:C5',columns={'need':'A','unmet':'B','partial':'C'}),
           dict(id='f',op='filter',input='r',where={'column':'need','cmp':'ne','value':'any'}),
           dict(id='d',op='derive',input='f',columns={'both':{'op':'add','args':[{'col':'unmet'},{'col':'partial'}]}}),
           dict(id='s',op='sort',input='d',by=[{'column':'both','descending':True,'type':'number'}]),
           dict(id='t',op='take',input='s',count=1),fmt={'columns':['both'],'decimals':1,'percent':True})
    text,trace=execute(p,shares,True)
    assert text=='50.0%' and trace[-1]['percent_corrected']==['both']
    assert execute(p,shares,False)[0]=='50.0%'


def test_a_ratio_of_amounts_asked_as_a_percentage_is_multiplied(shares):
    p=plan(dict(id='r',op='read',sheet=0,range='A3:E5',columns={'need':'A','men':'D','women':'E'}),
           dict(id='f',op='filter',input='r',where={'column':'need','cmp':'eq','value':'counselling'}),
           dict(id='d',op='derive',input='f',columns={'ratio':{'op':'div','args':[{'col':'women'},{'col':'men'}]}}),
           fmt={'columns':['ratio'],'decimals':1})
    text,trace=execute(p,shares,True)
    assert text=='81.3' and trace[-1]['percent_corrected']==['ratio']
    assert execute(p,shares,False)[0]=='0.8'
    p['steps'][2]['columns']['ratio']={'op':'mul','args':[{'literal':100},p['steps'][2]['columns']['ratio']]}
    p['format']['percent']=True
    assert execute(p,shares,True)[0]=='81.3%'


def test_percent_formatted_fractions_follow_the_question(shares):
    p=plan(dict(id='r',op='read',sheet=0,range='A3:G5',columns={'need':'A','share':'F','bare':'G'}),
           dict(id='a',op='aggregate',input='r',metrics={'top':{'fn':'max','column':'share'},'low':{'fn':'min','column':'share'}}),
           dict(id='d',op='derive',input='a',columns={'gap':{'op':'sub','args':[{'col':'top'},{'col':'low'}]}}),
           fmt={'columns':['gap'],'decimals':1})
    assert execute(p,shares,True)[0]=='37.5'
    assert execute(p,shares,False)[0]=='0.4'
    # Unformatted numbers within -2..2 may be shares or small percent figures: the plan's flag decides.
    p['steps'][1]['metrics']={'top':{'fn':'max','column':'bare'},'low':{'fn':'min','column':'bare'}}
    assert execute(p,shares,True)[0]=='0.4'
    p['format']['percent']=True
    assert execute(p,shares,True)[0]=='37.5%'


def test_count_shares_scale_but_means_keep_their_units(shares):
    p=plan(dict(id='r',op='read',sheet=0,range='A3:C5',columns={'need':'A','unmet':'B'}),
           dict(id='a',op='aggregate',input='r',metrics={'n':{'fn':'count'},'total':{'fn':'sum','column':'unmet'},
                'named':{'fn':'count','where':{'column':'need','cmp':'ne','value':'any'}}}),
           dict(id='d',op='derive',input='a',columns={'part':{'op':'div','args':[{'col':'named'},{'col':'n'}]},
                                                   'mean':{'op':'div','args':[{'col':'total'},{'col':'n'}]}}),
           fmt={'columns':['part','mean'],'decimals':1,'percent':True})
    assert execute(p,shares,True)[0]=='66.7%, 25.4%'


@pytest.mark.parametrize('query,asks',[
    ('Thu nhập của nữ bằng bao nhiêu phần trăm so với nam?',True),
    ('What percentage of the units are wood?',True),
    ("Quel est l'écart (en points de pourcentage) entre les groupes ?",True),
    ('몇 퍼센트 감소했습니까?',True),
    ('增长了百分之几？',True),
    ('Tỷ lệ (%) của nhóm A là bao nhiêu?',True),
    ('Answer with only the number followed by a % sign.',True),
    ('Doanh thu tháng 2 gấp bao nhiêu lần tháng 1?',False),
    ('What is the ratio of the percent change for Services to that for Commodities?',False),
    ('What is the 90th percentile of income?',False),
    ('Which year had an unemployment rate above 5%?',False),
])
def test_percentage_requests_are_told_from_ratio_questions(query,asks):
    assert planner.asks_percent(query) is asks


def test_planned_share_is_printed_as_the_requested_percentage(shares,monkeypatch):
    reply=plan(dict(id='r',op='read',sheet=0,range='A3:E5',columns={'need':'A','men':'D','women':'E'}),
               dict(id='f',op='filter',input='r',where={'column':'need','cmp':'eq','value':'counselling'}),
               dict(id='d',op='derive',input='f',columns={'ratio':{'op':'div','args':[{'col':'women'},{'col':'men'}]}}),
               fmt={'columns':['ratio'],'decimals':1})
    monkeypatch.setattr(L,'_post',lambda body,timeout:response(reply))
    cache=run.TableCache();b=L.Budget()
    query=('Thu nhập trung vị của nữ bằng bao nhiêu phần trăm so với nam? Làm tròn đến một chữ số thập phân, '
           'chỉ trả lời bằng số, không kèm ký hiệu %.')
    try: assert planner.answer(query,shares.path,cache,b,'raw','q')=='81.3'
    finally: planner.close_cache(cache)
    assert b.attempts==1 and b.diagnostics['trace'][-1]['percent_corrected']==['ratio']
