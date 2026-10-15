import sys
from pathlib import Path

import openpyxl
import pytest
from openpyxl.styles import Alignment

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from src.code_planner import plan
from src.plan_executor import execute
from src.xlsx_tables import load_tables


@pytest.fixture
def tables(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Category", "Percentage of deliveries", "Count"])
    for label, percent, count, indent in [
        ("Total", 100, 200, 0), ("Express", 70, 140, 1),
        ("Air", 40, 80, 2), ("Rail", 30, 60, 2),
        ("Standard", 30, 60, 1), ("Road", 20, 40, 2), ("Sea", 10, 20, 2),
    ]:
        ws.append([label, percent, count])
        ws.cell(ws.max_row, 1).alignment = Alignment(indent=indent)
    path = tmp_path / "deliveries.xlsx"
    wb.save(path)
    return load_tables(path)


@pytest.mark.parametrize("direction,expected", [("highest", "Express"), ("lowest", "Sea")])
def test_category_extremum_is_not_a_share_calculation(tables, direction, expected):
    query = f"Identify the category with {direction} percentage of deliveries. Answer with the category name."
    built = plan(query, tables)
    assert built["op"] == ("argmax" if direction == "highest" else "argmin")
    assert execute(built, tables, query) == expected


def test_individual_category_request_keeps_leaf_filter(tables):
    query = "Identify the individual category with highest percentage of deliveries. Answer with the category name."
    built = plan(query, tables)
    assert built["rows"]["exclude_subtotals"]
    assert execute(built, tables, query) == "Air"


def test_plain_percentage_question_remains_share(tables):
    query = "What percentage of deliveries used Air? Answer with the number."
    built = plan(query, tables)
    assert built["op"] == "share"
