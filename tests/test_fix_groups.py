"""Synthetic-table tests for the generalization fix groups (KSP1)."""
import openpyxl
import pytest

from src.planner import plan_with_confidence
from src.plan_executor import execute
from src.xlsx_tables import load_tables


def _ans(path, q):
    t = load_tables(str(path))
    p, _, _ = plan_with_confidence(q, t, strict=False)
    return execute(p, t, q)


@pytest.fixture
def sections(tmp_path):
    wb = openpyxl.Workbook(); ws = wb.active
    ws.append(["Status by age"]); ws.append([])
    ws.append(["Age", "Employed", "Unemployed"])
    data = {"TOTAL": [(30, 9), (50, 7), (40, 3)], "Men": [(10, 5), (35, 2), (15, 1)], "Women": [(20, 4), (15, 5), (25, 2)]}
    ages = ["16 to 24 years", "25 to 54 years", "55 years and over"]
    for sec, vals in data.items():
        ws.append([sec])
        for a, (e, u) in zip(ages, vals):
            ws.append([a, e, u])
    f = tmp_path / "sec.xlsx"; wb.save(f); return f


def test_section_scope_sum(sections):
    assert _ans(sections, "What is the combined number of unemployed men across all age groups? Answer with only the integer.") == "8"


def test_section_scope_argmax(sections):
    assert _ans(sections, "Which age group of women had the most employed? Answer with only the age group.") == "55 years and over"


def test_numeric_row_label_not_condition(sections):
    assert _ans(sections, "Among men, the number unemployed aged 55 years and over was how many? Answer with only the integer.") == "1"
    assert _ans(sections, "How many women aged 55 years and over were employed? Answer with only the integer.") == "25"


# ---- Vietnamese (with and without diacritics)
def test_vi_section_and_range(sections):
    assert _ans(sections, "Số người có việc làm của nữ từ 25 đến 54 tuổi là bao nhiêu? Chỉ trả lời bằng số nguyên.") == "15"
    assert _ans(sections, "So nguoi that nghiep cua nam tu 55 tuoi tro len la bao nhieu? Chi tra loi bang so nguyen.") == "1"


def test_vi_argmax_nodiac(sections):
    assert _ans(sections, "Nhom tuoi nao cua nu co nhieu nguoi that nghiep nhat? Chi tra loi ten.") == "25 to 54 years"


def test_vi_average_is_value(tmp_path):
    wb = openpyxl.Workbook(); ws = wb.active
    ws.append(["Name", "Score"])
    for n, s in [("A", 70), ("B", 80), ("C", 91)]:
        ws.append([n, s])
    f = tmp_path / "s.xlsx"; wb.save(f)
    assert _ans(f, "Điểm trung bình của tất cả mọi người là bao nhiêu? Chỉ trả lời bằng số, làm tròn đến 2 chữ số thập phân.") == "80.33"


# ---- top N / ties / thousands separator
@pytest.fixture
def people(tmp_path):
    wb = openpyxl.Workbook(); ws = wb.active
    ws.append(["First", "Last", "Score"])
    for r in [("An", "Ka", 70), ("Bo", "Lu", 95), ("Cu", "Mo", 88), ("Di", "No", 95), ("Em", "Op", 60)]:
        ws.append(list(r))
    f = tmp_path / "p.xlsx"; wb.save(f); return f


def test_top_n_list(people):
    assert _ans(people, "List the three highest scorers. Answer with the last names in descending order of score, separated by a comma and one space.") == "Lu, No, Mo"


def test_ties_share_top(people):
    assert _ans(people, "Which people share the top score? Answer with the last names in table order, separated by a comma and one space.") == "Lu, No"


def test_thousands_separator(tmp_path):
    wb = openpyxl.Workbook(); ws = wb.active
    ws.append(["Item", "Units"]); ws.append(["Alpha", 1234567]); ws.append(["Beta", 89])
    f = tmp_path / "u.xlsx"; wb.save(f)
    assert _ans(f, "How many units did Alpha sell? Answer with the number formatted with commas.") == "1,234,567"
    assert _ans(f, "How many units did Alpha sell? Answer with only the integer.") == "1234567"


# ---- cross-sheet / cross-column direction
@pytest.fixture
def years(tmp_path):
    wb = openpyxl.Workbook(); wb.remove(wb.active)
    for y, vals in (("2021", [100, 50, 70]), ("2022", [90, 80, 75])):
        ws = wb.create_sheet(y); ws.append(["Product", "Value"])
        for p, v in zip(["Tea", "Salt", "Corn"], vals):
            ws.append([p, v])
    f = tmp_path / "y.xlsx"; wb.save(f); return f


def test_grow_from_to(years):
    assert _ans(years, "Did Tea grow from 2021 to 2022? Answer with only Yes or No.") == "No"
    assert _ans(years, "Did Salt grow from 2021 to 2022? Answer with only Yes or No.") == "Yes"


def test_biggest_increase_sheets(years):
    assert _ans(years, "Which product had the biggest increase from 2021 to 2022? Answer with only the product name.") == "Salt"


def test_biggest_increase_columns(tmp_path):
    wb = openpyxl.Workbook(); ws = wb.active
    ws.append(["Region", "Sales 2021", "Sales 2022"])
    for r in [("North", 10, 30), ("South", 50, 55), ("East", 20, 25)]:
        ws.append(list(r))
    f = tmp_path / "c.xlsx"; wb.save(f)
    assert _ans(f, "Which region had the largest increase in sales from 2021 to 2022? Answer with only the region name.") == "North"


# ---- condition parser
@pytest.fixture
def recs(tmp_path):
    import datetime as dt
    wb = openpyxl.Workbook(); ws = wb.active
    ws.append(["Last", "Born", "Score", "Hours"])
    for r in [("Ka", dt.date(1970, 1, 2), 60, 10), ("Lu", dt.date(1982, 5, 6), 72, 40), ("Mo", dt.date(1990, 3, 3), 85, 20),
              ("No", dt.date(1985, 8, 9), 90, 45), ("Op", dt.date(1979, 1, 1), 72, 5)]:
        ws.append(list(r))
    f = tmp_path / "r.xlsx"; wb.save(f); return f


def test_between_inclusive(recs):
    assert _ans(recs, "How many people have a score between 72 and 85 inclusive? Answer with only the integer.") == "3"


def test_not_above(recs):
    assert _ans(recs, "Count the persons whose score is not above 72. Answer with only the integer.") == "3"


def test_born_after_and_at_least(recs):
    assert _ans(recs, "How many people born after 1980 scored at least 80? Answer with only the integer.") == "2"


def test_two_conditions_column_after_number(recs):
    assert _ans(recs, "How many people have more than 15 hours and more than 70 score? Answer with only the integer.") == "3"


def test_excluding_row(sections):
    assert _ans(sections, "Excluding the 16 to 24 group, what is the total employed among women? Answer with only the integer.") == "40"


def test_equal_to_ties(recs):
    assert _ans(recs, "Which people have a score equal to 72? Answer with the last names in table order, separated by a comma and one space.") == "Lu, Op"


# ---- passive / short / possessive forms
@pytest.fixture
def staff(tmp_path):
    import datetime as dt
    wb = openpyxl.Workbook(); ws = wb.active
    ws.append(["First", "Last", "Admission", "Score"])
    for r in [("An", "Ka", dt.date(2010, 1, 2), 60), ("Bo", "Lu", dt.date(2005, 5, 6), 72), ("An", "Mo", dt.date(2012, 3, 3), 85)]:
        ws.append(list(r))
    f = tmp_path / "st.xlsx"; wb.save(f); return f


def test_possessive_date(staff):
    assert _ans(staff, "Lu's admission date? Answer in YYYY-MM-DD format.") == "2005-05-06"


def test_passive_full_name(staff):
    assert _ans(staff, "The score obtained by An Mo was what? Answer with only the integer.") == "85"


def test_admitted_earliest(staff):
    assert _ans(staff, "Who was admitted earliest? Answer with only the last name.") == "Lu"


def test_by_how_much_exceed(sections):
    assert _ans(sections, "By how much did the employed count of the 25 to 54 group exceed that of the 16 to 24 group, for both sexes? Answer with only the integer.") == "20"
