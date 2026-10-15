"""Guards that turn confident wrong plans into a blank.

Two checks remain: answer type, and an nth rank from place 2 with no measure
column. Synthetic sheds and cargo only. Wording is written for these sheets.
"""
import os
import sys

import openpyxl

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))
from src.answer_type import answer_ok
from src.code_planner import plan, plan_with_confidence
from src.plan_executor import execute
from src.xlsx_tables import load_tables


def _answer(tables, query):
    planned = plan(query, tables)
    if not planned:
        return None
    return execute(planned, tables, query)


def _reasons(tables, query):
    _plan, _conf, reasons = plan_with_confidence(query, tables)
    return reasons


def _ledger(path):
    """Title row, two header levels, four sheds."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "Shed cargo ledger"
    ws.merge_cells("A1:E1")
    ws["A3"] = "Shed"
    ws.merge_cells("A3:A4")
    ws["B3"] = "Cargo"
    ws.merge_cells("B3:C3")
    ws["D3"] = "Fruit"
    ws.merge_cells("D3:E3")
    ws["B4"] = "Ripe"
    ws["C4"] = "Green"
    ws["D4"] = "Ripe"
    ws["E4"] = "Green"
    for i, (name, ripe, green, fruit, other) in enumerate(
        [("Amber", 40, 5, 8, 2), ("Birch", 25, 9, 30, 1),
         ("Cedar", 18, 3, 12, 7), ("Dune", 9, 4, 4, 6)],
        start=5,
    ):
        ws.cell(i, 1, name)
        ws.cell(i, 2, ripe)
        ws.cell(i, 3, green)
        ws.cell(i, 4, fruit)
        ws.cell(i, 5, other)
    wb.save(path)
    return load_tables(str(path))


def _stub_header(path):
    """A title, a real header, then a numeric stub row the reader keeps as data."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "Shed ledger title"
    ws.merge_cells("A1:C1")
    ws["A2"] = "Shed"
    ws["B2"] = "Ripe"
    ws["C2"] = "Green"
    ws["A3"] = "(header)"
    ws["B3"] = 1
    ws["C3"] = 2
    ws["A4"] = "Amber"
    ws["B4"] = 40
    ws["C4"] = 5
    ws["A5"] = "Birch"
    ws["B5"] = 25
    ws["C5"] = 9
    ws["A6"] = "Cedar"
    ws["B6"] = 18
    ws["C6"] = 4
    wb.save(path)
    tables = load_tables(str(path))
    assert tables[0].df["_row_path"].tolist()[0] == "(header)"
    return tables


def _orchard(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Band", "Pears", "Plums"])
    for row in [("Total, span 9 and over", 100, 80), ("span 1 to 3", 10, 12),
                ("span 4 to 6", 40, 30), ("span 7 to 8", 20, 22)]:
        ws.append(list(row))
    wb.save(path)
    return load_tables(str(path))


def _dual(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Tag", "Grams", "Lux"])
    for row in [("Nia", 88, 3), ("Pax", 63, 9), ("Quill", 52, 8), ("Ren", 41, 1)]:
        ws.append(list(row))
    wb.save(path)
    return load_tables(str(path))


def _grams(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Tag", "Grams"])
    for row in [("Nia", 88), ("Pax", 63), ("Quill", 52), ("Ren", 41)]:
        ws.append(list(row))
    wb.save(path)
    return load_tables(str(path))


def test_topn_list_is_not_a_numeric_aggregate(tmp_path):
    # Before: sum of the bushels column, '96'.
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Pier", "Bushels"])
    for row in [("Amber", 40), ("Birch", 25), ("Cedar", 18), ("Dune", 9), ("Elm", 4)]:
        ws.append(list(row))
    path = tmp_path / "piers.xlsx"
    wb.save(path)
    tables = load_tables(str(path))
    q = "Name three pier labels using the total bushels figure, ordered downward and comma-separated."
    assert _answer(tables, q) is None
    assert "blank:shape" in _reasons(tables, q)


def test_resolved_topn_on_multilevel_header_still_answers(tmp_path):
    tables = _ledger(tmp_path / "ledger.xlsx")
    q = ("Name the three sheds with the largest Cargo Ripe, descending, "
         "separated by a comma and one space.")
    assert _answer(tables, q) == "Amber, Birch, Cedar"


def test_complete_rank_is_not_second_guessed(tmp_path):
    # The question says 'ripe' and two Ripe columns exist. The planner still
    # emits a rank with N, a column, and a direction. That plan is kept.
    tables = _ledger(tmp_path / "ledger.xlsx")
    q = "Name three sheds with the largest ripe figure, comma-separated."
    assert _answer(tables, q) == "Amber, Birch, Cedar"
    assert not any(str(r).startswith("blank:") for r in _reasons(tables, q))


def test_complete_rank_keeps_both_direction_words(tmp_path):
    # The planner already chose N, a column, and an order. Guards do not override it.
    tables = _ledger(tmp_path / "ledger.xlsx")
    q = "Name the three largest Cargo Ripe sheds, ascending, separated by a comma."
    assert not any(str(r).startswith("blank:") for r in _reasons(tables, q))
    assert _answer(tables, q) == "Dune, Cedar, Birch"


def test_vietnamese_topn_with_column_and_direction_answers(tmp_path):
    tables = _grams(tmp_path / "grams.xlsx")
    assert _answer(tables, "Grams: ba lớn nhất.") == "Nia, Pax, Quill"
    assert _answer(tables, "Grams: ba lon nhat.") == "Nia, Pax, Quill"


def test_scalar_extreme_stays_a_single_number(tmp_path):
    tables = _grams(tmp_path / "grams.xlsx")
    assert _answer(tables, "highest 3 Grams. Respond with an integer.") == "88"
    assert _answer(tables, "top 2 Grams. Respond with an integer.") == "88"


def test_shorter_topn_list_stays_when_the_plan_is_complete(tmp_path):
    # A filtered or short table can legitimately return fewer than N labels.
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Shed", "Ripe"])
    ws.append(["Amber", 40])
    ws.append(["Birch", 25])
    path = tmp_path / "short.xlsx"
    wb.save(path)
    tables = load_tables(str(path))
    q = "Name the three sheds with the largest Ripe, descending, separated by a comma and one space."
    assert _answer(tables, q) == "Amber, Birch"
    planned, conf, reasons = plan_with_confidence(q, tables)
    assert planned and planned["op"] == "rank" and conf >= 0.5
    assert not any(str(r).startswith("blank:") for r in reasons)


def test_nth_without_a_rank_column_blanks(tmp_path):
    # Before: an arbitrary tag, 'Pax', confidence 1.0, by=None.
    tables = _dual(tmp_path / "dual.xlsx")
    q = "Which tag is the second highest?"
    assert _answer(tables, q) is None
    assert "blank:nth-column" in _reasons(tables, q)
    assert _answer(tables, "The 2nd-largest tag.") is None


def test_nth_with_one_measure_still_answers(tmp_path):
    tables = _grams(tmp_path / "grams.xlsx")
    assert _answer(tables, "Which tag ranks second when Grams decides the order?") == "Pax"
    assert _answer(tables, "Find the 2nd-largest Grams tag.") == "Pax"


def test_nth_names_its_measure_on_a_two_column_sheet(tmp_path):
    # Place, measure, and direction are already on the plan. The other measure is ignored.
    tables = _dual(tmp_path / "dual.xlsx")
    q = "Which tag is the second highest Grams?"
    assert not any(str(r).startswith("blank:") for r in _reasons(tables, q))
    assert _answer(tables, q) == "Pax"


def test_ordinal_inside_a_column_name_still_answers(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Region", "2nd quarter", "3rd quarter"])
    ws.append(["North", 11, 15])
    ws.append(["South", 7, 9])
    path = tmp_path / "quarters.xlsx"
    wb.save(path)
    tables = load_tables(str(path))
    q = "What were 2nd quarter arrivals in North? Answer with only the integer."
    assert _answer(tables, q) == "11"


def test_pair_single_number_blanks(tmp_path):
    # Before: one summed figure, '134'.
    tables = _orchard(tmp_path / "orch.xlsx")
    q = ("Pears and plums for the total older span, in that order, "
         "separated by a comma and one space.")
    assert _answer(tables, q) is None
    assert "blank:shape" in _reasons(tables, q)
    assert not answer_ok("number", "134", {"kind": "list", "count": 2, "sep": ", "}, q, {"op": "sum"})


def test_pair_of_cells_still_answers(tmp_path):
    tables = _orchard(tmp_path / "orch.xlsx")
    q = ("What are the pears and plums of span 4 to 6? "
         "Answer as two integers in the order `pears, plums`, separated by a comma and one space.")
    assert _answer(tables, q) == "40, 30"


def test_combined_sum_still_answers(tmp_path):
    tables = _orchard(tmp_path / "orch.xlsx")
    q = "What is the sum of pears and plums across all spans? Answer with only the integer."
    assert _answer(tables, q) == "134"
    tables = _ledger(tmp_path / "ledger.xlsx")
    q = "What is the sum of Cargo Ripe across all sheds? Answer with only the integer."
    assert _answer(tables, q) == "92"


def test_total_asks_for_a_sum(tmp_path):
    # 'total' here asks for one summed integer. The planner and the executor both run.
    tables = _ledger(tmp_path / "ledger.xlsx")
    q = "What is the total of Cargo Ripe across every shed? Answer with only the integer."
    assert not any(str(r).startswith("blank:") for r in _reasons(tables, q))
    assert _answer(tables, q) == "92"


def test_named_row_lookup_still_answers(tmp_path):
    tables = _ledger(tmp_path / "ledger.xlsx")
    assert _answer(tables, "Read the Cargo Ripe cell.") == "40"
    q = "Cargo Ripe recorded against Birch. Answer with only the integer."
    assert _answer(tables, q) == "25"


def test_lookup_on_a_single_row_still_answers(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Shed", "Ripe"])
    ws.append(["Only", 7])
    path = tmp_path / "one.xlsx"
    wb.save(path)
    tables = load_tables(str(path))
    assert _answer(tables, "Read the Ripe cell.") == "7"


def test_header_stub_lookup_follows_the_reader(tmp_path):
    # The reader keeps the numeric stub as a data row. A named row still resolves.
    tables = _stub_header(tmp_path / "stub.xlsx")
    assert _answer(tables, "Read the Ripe cell.") == "1"
    assert _answer(tables, "Ripe recorded against Amber. Answer with only the integer.") == "40"


def test_single_number_question_rejects_a_label_list():
    fmt = {"kind": "text", "sep": ", "}
    q = "Tong so bushels la bao nhieu?"
    assert not answer_ok("any", "Amber, Birch", fmt, q, {"op": "rank"})
    assert answer_ok("number", "96", {"kind": "int", "sep": ", "},
                     "What is the sum of bushels? Answer with only the integer.",
                     {"op": "sum"})


def _kiosks(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Kiosk", "Bulbs", "Cords"])
    for row in [("Maple", 12, 4), ("Nook", 30, 2), ("Orbit", 7, 9), ("Pylon", 18, 1)]:
        ws.append(list(row))
    wb.save(path)
    return load_tables(str(path))


def test_single_superlative_with_total_stays_a_label(tmp_path):
    tables = _kiosks(tmp_path / "kiosks.xlsx")
    q = "Name the kiosk with the highest total bulbs."
    assert _answer(tables, q) == "Nook"
    reasons = _reasons(tables, q)
    assert "blank:shape" not in reasons
    assert "blank:nth-column" not in reasons


def test_guards_are_noop_when_rank_column_and_direction_are_set():
    from src.planner import _blank_reasons
    # Hostile wording: both directions, an ordinal, and the word total.
    q = "four largest bulbs ascending, highest first, total column, second highest"
    plan = {
        "op": "rank",
        "args": {"k": 4, "by": "Bulbs", "order": "asc"},
        "columns": ["Bulbs"],
        "format": {"kind": "list", "sep": ", "},
        "rows": {"paths": [], "where": []},
    }
    assert _blank_reasons(q, q, plan, {}) == []
    # N is not required. Column and direction are enough for a rank plan.
    no_n = {
        "op": "rank",
        "args": {"by": "Bulbs", "order": "desc"},
        "columns": [],
        "format": {"kind": "list", "sep": ", "},
        "rows": {"paths": [], "where": []},
    }
    assert _blank_reasons(q, q, no_n, {}) == []
    nth = {
        "op": "nth",
        "args": {"n": 2, "by": "Grams", "order": "desc", "return": "Tag"},
        "columns": ["Grams", "Lux"],
        "format": {"kind": "text"},
        "rows": {"paths": [], "where": []},
    }
    qn = "Which tag is the second highest?"
    assert _blank_reasons(qn, qn, nth, {}) == []
    one = {
        "op": "argmax",
        "args": {"ties": "first", "over": "rows"},
        "columns": ["Bulbs"],
        "format": {"kind": "text"},
        "rows": {"paths": [], "where": []},
    }
    q1 = "Name the kiosk with the highest total bulbs."
    assert _blank_reasons(q1, q1, one, {}) == []


def test_rank_with_a_clear_count_is_not_blanked(tmp_path):
    tables = _kiosks(tmp_path / "kiosks.xlsx")
    q = "List the 3 kiosks with the largest bulb count, separated by a comma and one space."
    assert not any(str(r).startswith("blank:") for r in _reasons(tables, q))
    assert _answer(tables, q) == "Nook, Pylon, Maple"
