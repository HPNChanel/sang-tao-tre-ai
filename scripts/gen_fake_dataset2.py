#!/usr/bin/env python3
"""Write a small synthetic dataset-2 pack for the dry run.

Pure Python, openpyxl, fixed seed. Not contest data. The pack matches the
layout run.py reads:

  <out>/queries.json
  <out>/sample_submission.csv
  <out>/xlsx_files/dataset2_fake.xlsx

The workbook has three kinds of sheet, same shapes as round 1:
  People  — record table, two-row merged header
  Stats   — grouped rows plus a Total row
  2023, 2024, 2025 — same columns, one sheet per period
"""
from __future__ import annotations

import argparse
import csv
import json
import random
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font

SEED = 20261016
XLSX_NAME = "dataset2_fake.xlsx"


def _people(wb, rng: random.Random):
    ws = wb.active
    ws.title = "People"
    ws["A1"] = "No"
    ws.merge_cells("A1:A2")
    ws["B1"] = "Name"
    ws.merge_cells("B1:C1")
    ws["B2"] = "First"
    ws["C2"] = "Last"
    ws["D1"] = "Born"
    ws.merge_cells("D1:D2")
    ws["E1"] = "Score"
    ws.merge_cells("E1:E2")
    for col in ("A", "B", "D", "E"):
        ws[f"{col}1"].alignment = Alignment(horizontal="center", vertical="center")
    ws["B1"].alignment = Alignment(horizontal="center")
    first = ["An", "Binh", "Chi", "Dung", "Ha"]
    last = ["Le", "Tran", "Pham", "Vo", "Do"]
    rng.shuffle(last)
    for i, (fn, ln) in enumerate(zip(first, last), start=1):
        row = i + 2
        year = rng.randint(1988, 2004)
        month = rng.randint(1, 12)
        day = rng.randint(1, 28)
        score = rng.choice([45, 60, 70, 80, 90])
        ws.cell(row, 1, i)
        ws.cell(row, 2, fn)
        ws.cell(row, 3, ln)
        ws.cell(row, 4, datetime(year, month, day))
        ws.cell(row, 4).number_format = "YYYY-MM-DD"
        ws.cell(row, 5, score)


def _stats(wb):
    ws = wb.create_sheet("Stats")
    ws["A1"] = "Synthetic employment counts"
    ws.merge_cells("A1:C1")
    ws["A1"].font = Font(bold=True)
    ws["A3"] = "Group"
    ws["B3"] = "Employed"
    ws["C3"] = "Unemployed"
    # Counts are fixed so a Total row is the sum of the four group rows.
    rows = [
        (4, "Total", 100, 17),
        (5, "Men", None, None),
        (6, "16 to 24", 10, 5),
        (7, "25 to 54", 20, 1),
        (8, "Women", None, None),
        (9, "16 to 24", 30, 2),
        (10, "25 to 54", 40, 9),
    ]
    for r, label, employed, unemployed in rows:
        ws.cell(r, 1, label)
        if employed is None:
            continue
        ws.cell(r, 2, employed)
        ws.cell(r, 3, unemployed)


def _periods(wb, rng: random.Random):
    regions = ["North", "South", "East"]
    for year in (2023, 2024, 2025):
        ws = wb.create_sheet(str(year))
        ws["A1"] = "Region"
        ws["B1"] = "Sales"
        for i, region in enumerate(regions):
            ws.cell(i + 2, 1, region)
            ws.cell(i + 2, 2, rng.randint(10, 80))


def _queries():
    fname = XLSX_NAME
    items = [
        (
            "d2q01",
            "What score did the oldest person receive? Answer with only the integer score.",
        ),
        (
            "d2q02",
            "How many people scored at least 60? Answer with only the integer.",
        ),
        (
            "d2q03",
            "How many women aged 25 to 54 were employed? Answer with only the integer.",
        ),
        (
            "d2q04",
            "Which age group had the most unemployed women? Answer with only the age group.",
        ),
        (
            "d2q05",
            "Which region had the highest Sales in 2024? Answer with only the region name.",
        ),
    ]
    return [{"id": i, "file_name": fname, "query": q} for i, q in items]


def write_pack(out: Path) -> None:
    out.mkdir(parents=True, exist_ok=True)
    xlsx_dir = out / "xlsx_files"
    xlsx_dir.mkdir(parents=True, exist_ok=True)
    rng = random.Random(SEED)
    wb = openpyxl.Workbook()
    _people(wb, rng)
    _stats(wb)
    _periods(wb, rng)
    wb.save(xlsx_dir / XLSX_NAME)
    queries = _queries()
    (out / "queries.json").write_text(
        json.dumps(queries, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    # Id order differs from queries.json so a sample file is actually exercised.
    sample_ids = ["d2q05", "d2q01", "d2q03", "d2q02", "d2q04"]
    with open(out / "sample_submission.csv", "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(["id", "answer"])
        for qid in sample_ids:
            w.writerow([qid, ""])


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description="Write a synthetic dataset-2 pack.")
    ap.add_argument("--out", required=True, help="Directory to create")
    args = ap.parse_args(argv)
    write_pack(Path(args.out))
    print(args.out)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
